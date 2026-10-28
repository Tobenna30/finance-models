from pathlib import Path
import importlib.util,sys
import numpy as np
import openpyxl
import pytest
P=Path(__file__).resolve().parents[1]/'projects/12-rates-macro/policy'

def load():
 old=sys.modules.pop('assumptions',None);sys.path.insert(0,str(P))
 spec=importlib.util.spec_from_file_location('rates_policy_engine',P/'engine.py');e=importlib.util.module_from_spec(spec);spec.loader.exec_module(e);sys.path.pop(0);sys.modules.pop('assumptions',None)
 if old:sys.modules['assumptions']=old
 return e
E=load();R=E.results()

def test_source_dates_and_no_duplicates():
 d=E.curves();assert d.index.is_unique;assert d.index.max()==E.A.ASOF;assert not d.isna().any().any();assert len(d)>1500

def test_policy_intervals_reproduce_discount_integral():
 for m in E.meeting_path():
  daily=(m['market']-E.A.BASIS_BP/100)/100
  cc=365*np.log1p(daily/365)
  assert cc*(m['t2']-m['t1'])==pytest.approx(E.integral(m['t2'])-E.integral(m['t1']),abs=1e-12)

def test_gilt_cashflows_and_parallel_neutrality():
 b=R['bonds'];assert all(x['cf'][-1][1]>100 for x in b)
 assert sum(x['nominal_gbp']/100*x['dv01_per100'] for x in b)==pytest.approx(0,abs=1e-7)
 assert sum(abs(x['nominal_gbp']/100*x['dv01_per100']) for x in b)==pytest.approx(E.A.BUDGET)

def test_carry_roll_cash_reconciles():
 for b in R['bonds']:
  raw=E.bond_price(b['cf'],elapsed=E.A.HORIZON)+sum(c for t,c in b['cf'] if t<=E.A.HORIZON)-b['price']-b['price']*E.A.BANK_RATE/100*E.A.HORIZON
  assert raw==pytest.approx(b['carry']+b['roll'])

def test_hac_ols_known_slope():
 x=np.arange(40);y=3+2*x+np.sin(x)
 v=E.ols(x,y);assert v['slope']==pytest.approx(2,abs=.01);assert v['t']>100;assert v['r2']>.99

def test_fx_quote_and_no_lookahead():
 d=E.fx_panel();assert d.index.max()<=E.A.ASOF
 assert d.GBP_EUR.std()>0;assert R['fx']['GBP_USD']['regressions']['ex2022']['n']<len(d)
 assert R['fx']['GBP_EUR']['regressions']['CPI']['n']==9

def test_replay_is_full_cashflow_reprice():
 v=R['replay']['peak'];expected=sum(b['nominal_gbp']/100*(E.bond_price(b['cf'],s/10000)-b['price']) for b,s in zip(R['bonds'],v['shocks_bp']))
 assert v['pnl_gbp']==pytest.approx(expected)

def test_excel_ties_and_no_errors():
 p=P.parent/'models/Gilt_GBP_Policy.xlsx';w=openpyxl.load_workbook(p,data_only=True)
 for s in w:
  assert not [(c.coordinate,c.value) for row in s for c in row if c.data_type=='e']
 for cell,key in [('B2','implied_bank_rate_12m_pct'),('B3','implied_cuts_bp'),('B5','spread_bp'),('B7','stop_bp'),('B14','carry_roll_bp')]:assert w['Trade'][cell].value==pytest.approx(R[key],abs=1e-6)
 assert w['Trade']['B15'].value==pytest.approx(0,abs=1e-6)
 f=openpyxl.load_workbook(p,data_only=False);assert f['Trade']['D10'].data_type=='f';assert f['Policy path']['H2'].data_type=='f'
 assert [w['Sensitivity'].cell(i,4).value for i in range(2,7)]==sorted(w['Sensitivity'].cell(i,4).value for i in range(2,7))
