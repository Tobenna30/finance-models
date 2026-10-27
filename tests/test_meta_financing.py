"""Issuer funding checks, initial conversion boundaries and native Excel edits."""
import importlib
import json
import os
import shutil
import subprocess
import tempfile
from pathlib import Path
import pytest
from openpyxl import load_workbook

P='projects.06-ecm-dcm-financing.meta.'
engine=importlib.import_module(P+'engine')
a=importlib.import_module(P+'assumptions')
b=importlib.import_module(P+'build_excel')

def test_primary_facts_and_funding_timing():
    r=engine.calculate()
    assert r['gross_package']==37500
    assert r['initial_cash']==12475
    assert r['forward_net']==pytest.approx(40*622.92-10)
    assert r['net_total']==pytest.approx(37381.8)
    assert r['initial_cash'] < r['net_total']/2
    assert r['historical_fcf']==11049
    assert r['liquid_resources']==82449
    assert r['capex_remaining_low']==49176
    assert r['capex_remaining_high']==59176

def test_depositary_and_preferred_not_confused():
    r=engine.calculate()
    assert r['pref_count']==12.5
    assert r['annual_dividend']==812.5
    assert r['conversion_shares']==20
    assert r['total_new_initial_term_shares']==60

@pytest.mark.parametrize('price,rate',[(400,1.6),(625,1.6),(700,1.4286),(1000/1.362,1.362),(850,1.362)])
def test_initial_conversion_boundaries(price,rate):
    r=engine.calculate({'price_base':price})
    assert r['conversion_rate']==rate
    assert r['conversion_shares']==pytest.approx(12.5*rate)

def test_security_return_order_and_hurdle():
    r=[engine.calculate({'case':i}) for i in (1,2,3)]
    assert r[0]['pref_total_return'] < r[1]['pref_total_return'] < r[2]['pref_total_return']
    z=r[1];price=z['investor_hurdle_price']
    at=engine.calculate({'price_base':price})
    pv=at['dividend_pv']+at['converted_value']/1.08**((a.CONVERT-a.CLOSE).days/365)
    assert pv==pytest.approx(12500,abs=1e-7)
    assert at['investor_hurdle_price']>734

def test_bridge_and_net_proceeds_reconcile():
    r=engine.calculate();v=a.values()
    assert r['bridge_after_tax_cost']==pytest.approx((24906.8*.055*(a.SETTLE-a.CLOSE).days/365+24906.8*.005)*.79)
    assert r['bond_face']*(1-v['bond_fee'])==pytest.approx(r['net_total'])
    assert r['equity_gross']*(1-v['equity_fee'])==pytest.approx(r['net_total'])
    assert engine.calculate({'forward_change':-.05})['net_total'] < r['net_total']

def test_effective_interest_and_bullet():
    r=engine.calculate();s=r['bond_schedule']
    assert len(s)==5 and s[-1]['closing']==pytest.approx(0,abs=1e-7)
    assert sum(y['amortisation'] for y in s)==pytest.approx(r['bond_face']-r['net_total'])
    assert sum(y['principal'] for y in s)==r['bond_face']
    assert r['bond_effective_yield']>.055
    assert r['bond_price']==pytest.approx(100)
    assert engine.calculate({'bond_yield':.065})['bond_price']<100
    for y in s:assert y['opening']+y['amortisation']-y['principal']==pytest.approx(y['closing'])

def test_dividend_stub_and_position_balances():
    r=engine.calculate()
    assert len(r['dividends'])==11
    assert r['dividends'][0]['amount']==pytest.approx(12500*.065*106/360)
    assert sum(x['amount'] for x in r['dividends'])==pytest.approx(2270.4861111111113)
    for p in r['positions'].values():assert p['assets']==pytest.approx(p['liabilities']+p['equity'])
    assert r['positions']['package_close']['cash']==23697
    assert r['positions']['package_settled']['cash']-r['positions']['package_close']['cash']==pytest.approx(r['forward_net'])

def parity(w,r):
    for row,k in {6:'pref_gross',8:'pref_net',9:'common_gross',10:'gross_package',14:'forward_net',15:'net_total',16:'initial_cash',17:'deferred_share',19:'annual_dividend',20:'bridge_interest',21:'bridge_fee',22:'bridge_after_tax_cost',24:'historical_fcf',25:'liquid_resources',26:'capex_remaining_low',27:'capex_remaining_high'}.items():
        assert w['Funding'][f'D{row}'].value==pytest.approx(r[k],abs=1e-7)
    for row,k in {6:'pref_count',7:'min_price',8:'max_price',9:'conversion_rate',10:'conversion_shares',11:'converted_value',12:'dividend_total',13:'pref_total_return',45:'dividend_pv',46:'required_terminal_value',47:'investor_hurdle_price'}.items():
        assert w['Conversion'][f'D{row}'].value==pytest.approx(r[k],abs=1e-6)
    for row,k in {6:'bond_face',7:'bond_coupon_cash',8:'bond_effective_yield',9:'bond_price'}.items():assert w['Bond'][f'D{row}'].value==pytest.approx(r[k],abs=1e-7)
    for i,y in enumerate(r['bond_schedule'],16):
        for col,key in zip('DEFGHI',('opening','coupon','interest','amortisation','principal','closing')):assert w['Bond'][f'{col}{i}'].value==pytest.approx(y[key],abs=1e-6)
    for col,key in zip('DEFG',('package_close','package_settled','immediate_equity','straight_debt')):
        for row,k in {8:'cash',9:'assets',13:'liabilities',14:'equity'}.items():assert w['Position'][f'{col}{row}'].value==pytest.approx(r['positions'][key][k],abs=1e-6)
        assert w['Position'][f'{col}11'].value==pytest.approx(0,abs=1e-7)

def test_cached_formulas_and_parity():
    w=load_workbook(b.OUT,data_only=True);parity(w,engine.calculate())
    assert not [(s.title,c.coordinate) for s in w for row in s for c in row if c.data_type=='e']
    f=load_workbook(b.OUT,data_only=False)
    for row in range(16,21):assert all(f['Bond'][f'{c}{row}'].data_type=='f' for c in 'DEFGHI')
    for row in range(20,31):assert f['Conversion'][f'E{row}'].data_type=='f'
    for key,v in a.INPUTS.items():assert f['Inputs'][f'D{b.I[key]}'].value==v[1] and v[2] in ('S','A','D') and v[4]
    saved=json.loads((a.ROOT/'outputs/results.json').read_text())
    for i,key in enumerate(('down','base','up'),1):assert saved[key]['pref_total_return']==pytest.approx(engine.calculate({'case':i})['pref_total_return'])

@pytest.mark.parametrize('change',[{'case':1},{'case':3},{'forward_change':-.05,'bond_yield':.065,'pref_fee':.003},{'deposit_price':48,'investor_hurdle':.06,'bond_coupon':.06}])
def test_native_input_changes(change):
    soffice=os.environ.get('SOFFICE') or shutil.which('soffice')
    if not soffice:pytest.skip('LibreOffice unavailable; cached parity remains mandatory')
    with tempfile.TemporaryDirectory() as tmp:
        tmp=Path(tmp);out=tmp/'out';out.mkdir();w=load_workbook(b.OUT)
        for key,val in change.items():w['Inputs'][f'D{b.I[key]}']=val
        f=tmp/'changed.xlsx';w.save(f)
        subprocess.run([soffice,'-env:UserInstallation='+(tmp/'profile').as_uri(),'--headless','--convert-to','xlsx:Calc MS Excel 2007 XML','--outdir',str(out),str(f)],capture_output=True,check=True,timeout=120)
        parity(load_workbook(out/f.name,data_only=True),engine.calculate(change))

def test_invalid_inputs():
    for change in ({'case':0},{'bond_term':6},{'bond_fee':1},{'price_base':0},{'forward_change':-1}):
        with pytest.raises(ValueError):engine.calculate(change)
