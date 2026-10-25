import sys
from pathlib import Path
import pytest
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "projects" / "04-08-darktrace"))

from darktrace import assumptions as A  # noqa: E402
from darktrace import engine as E  # noqa: E402

XLSX = Path(__file__).resolve().parents[1] / "projects" / "04-08-darktrace" / "models" / "Darktrace_04_08.xlsx"


def test_offer_maths_ties_to_announcement():
    D = A.DEAL
    assert abs(D["enterprise_value"] / D["ltm_revenue"] - D["ev_rev_stated"]) < 0.05
    assert abs(D["enterprise_value"] / D["ltm_ebitda"] - D["ev_ebitda_stated"]) < 0.05
    assert abs(D["offer_p"] / D["undisturbed_p"] - 1 - D["prem_close"]) < 0.005
    assert abs(D["offer_p"] / D["vwap3m_p"] - 1 - D["prem_vwap3m"]) < 0.005
    assert abs(A.NET_CASH - (A.H1["cash_dec23"] - A.H1["leases_dec23"])) < 0.1      # EV bridge = cash less leases
    ltm = A.HIST["Revenue"][2023] + A.H1["rev_h1fy24"] - A.H1["rev_h1fy23"]
    assert abs(ltm - D["ltm_revenue"]) < 1.0
    assert abs(A.FX - 1.25) < 1e-9


def test_sources_complete():
    keys = [c[-1] for c in A.COMPS] + [p[-2] for p in A.PRECEDENTS] + [p[-1] for p in A.PRECEDENTS]
    assert all(k in A.SOURCES for k in keys)
    assert all(u.startswith("https://") for _, u in A.SOURCES.values())


def test_dcf_consistency():
    d = E.dcf()
    assert abs(d["ev"] - (d["pv_fcf"] + d["pv_tv"])) < 1e-9
    g = E.dcf_grid()
    assert g[0][2] > g[2][2] > g[4][2]          # value falls as WACC rises
    assert g[2][0] < g[2][2] < g[2][4]          # and rises with terminal growth
    assert abs(g[2][2] - d["per_share_p"]) < 1e-9
    assert E.dcf(sbp_cost=False)["per_share_p"] > d["per_share_p"]


def test_lbo_sources_equal_uses_and_schedule_identities():
    for case in A.CASES:
        r = E.lbo(case)
        assert abs(r["sponsor"] + r["tlb0"] + r["sl0"] + r["cash_used"] - r["uses"]) < 1e-9
        tlb = r["tlb0"]
        for row in r["rows"]:
            assert abs(row["tlb"] - (tlb - row["amort"] - row["sweep"])) < 1e-9
            assert row["cash"] > 0 and row["sweep"] >= 0 and row["tax"] >= 0
            assert row["deductible"] <= row["interest"] + 1e-9
            tlb = row["tlb"]
    assert E.lbo("Base")["irr"] > E.lbo("Downside")["irr"]


def test_grid_and_ability_to_pay():
    g = E.lbo_grid()
    for row in g: assert all(a < b for a, b in zip(row, row[1:]))          # more debt -> higher IRR (base case)
    for col in zip(*g): assert all(a < b for a, b in zip(col, col[1:]))    # higher exit multiple -> higher IRR
    for hurdle in A.IRR_HURDLES:
        px = E.ability_to_pay(hurdle)
        assert abs(E.lbo(price_usd=px)["irr"] - hurdle) < 1e-9             # closed form is exact
    assert E.ability_to_pay(0.25) < E.ability_to_pay(0.20)


def test_football_field_brackets():
    for name, lo, hi in E.football():
        assert lo < hi, name


@pytest.fixture(scope="module")
def wb_values():
    from openpyxl import load_workbook
    return load_workbook(XLSX, data_only=True), load_workbook(XLSX)


def test_excel_ties_to_engine(wb_values):
    wv, _ = wb_values
    cover = {r[0]: r[1] for r in wv["Cover"].iter_rows(min_row=5, values_only=True) if r[0]}
    s = E.summary(); b = E.lbo("Base"); dn = E.lbo("Downside")
    checks = {"DCF value per share (p)": s["dcf_p"], "Peer median EV/NTM value (p)": s["comps_med_p"],
              "Precedent median EV/LTM value (p)": s["prec_med_p"], "LBO ability to pay at 20% IRR (p)": s["atp20_p"],
              "Strategic maximum (p)": s["strategic_p"], "LBO base IRR": b["irr"], "LBO base MOIC": b["moic"],
              "LBO downside IRR": dn["irr"], "LBO downside MOIC": dn["moic"], "Sponsor equity cheque ($m)": b["sponsor"],
              "Offer EV / NTM revenue": s["offer_ev_ntm"]}
    for k, v in checks.items():
        assert abs(cover[k] - v) < 1e-6 * max(1, abs(v)), k
    ff = [r for r in wv["Football field"].iter_rows(min_row=5, max_row=11, values_only=True)]
    for (name, lo, hi), row in zip(E.football(), ff):
        assert abs(row[1] - lo) < 1e-6 and abs(row[2] - hi) < 1e-6, name
    grid = E.lbo_grid()
    for i in range(len(A.GRID_EXIT)):
        for j in range(len(A.GRID_DEBT)):
            assert abs(wv["LBO grid"].cell(6 + i, 2 + j).value - grid[i][j]) < 1e-9


def test_excel_forecasts_are_formulas(wb_values):
    _, wf = wb_values
    for sheet, rng in (("DCF", "C5:I18"), ("LBO", "C21:G46"), ("Downside", "C21:G46")):
        vals = [c.value for row in wf[sheet][rng] for c in row if c.value is not None]
        assert vals and all(isinstance(v, str) and v.startswith("=") for v in vals), sheet


def test_value_bridge_sums(wb_values):
    steps, proceeds = E.bridge()
    assert abs(sum(v for _, v in steps) - proceeds) < 1e-6
    wv, _ = wb_values
    cover = {r[0]: r[1] for r in wv["Cover"].iter_rows(min_row=5, values_only=True) if r[0]}
    assert abs(cover["Bridge check (should be 0)"]) < 1e-6
