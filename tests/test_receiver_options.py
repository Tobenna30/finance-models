"""Saved Excel parity plus independent identities and native input-change tests."""
import importlib.util,sys,math,os,subprocess,zipfile,tempfile
from pathlib import Path
import xml.etree.ElementTree as ET
import pytest
import openpyxl

ROOT=Path(__file__).resolve().parents[1]
MOD=ROOT/'projects/12-rates-macro/options'
spec=importlib.util.spec_from_file_location('receiver_options',MOD/'__init__.py',submodule_search_locations=[str(MOD)])
pkg=importlib.util.module_from_spec(spec);sys.modules[spec.name]=pkg;spec.loader.exec_module(pkg)
from receiver_options import engine as E
BOOK=ROOT/'projects/12-rates-macro/models/Sterling_Receiver_Options.xlsx'
MAPPING={'forward':('Pricing','B24'),'annuity':('Pricing','B22'),'pv01':('Pricing','B23'),
 'outright_bp':('Summary','B13'),'spread_bp':('Summary','B14'),'spread_debit':('Summary','B15'),
 'spread_max_profit':('Summary','B16'),'spread_breakeven_bp':('Summary','B17'),'crossover_bp':('Summary','B18'),
 'linear_pnl':('Summary','B21'),'outright_pnl':('Summary','B22'),'spread_pnl':('Summary','B23'),
 'outright_delta':('Pricing','B38'),'spread_delta':('Pricing','B39'),
 'outright_vega':('Pricing','B40'),'spread_vega':('Pricing','B41')}
def parity(path,**updates):
 w=openpyxl.load_workbook(path,data_only=True);r=E.calculate(**updates)
 for key,(sn,cell) in MAPPING.items():assert w[sn][cell].value==pytest.approx(r[key],rel=1e-8,abs=1e-7),(key,w[sn][cell].value,r[key])
 return w
def test_saved_excel_parity():
 w=parity(BOOK)
 for s in w:
  for row in s:
   for c in row:assert c.data_type!='e',(s.title,c.coordinate,c.value)
 assert w['Summary']['B7'].value=='ILLUSTRATIVE VOLATILITY'
def test_formulas_and_source_separation():
 w=openpyxl.load_workbook(BOOK,data_only=False)
 for sn,cell in MAPPING.values():assert w[sn][cell].data_type=='f'
 for cell in ['B14','B16','B17','B18','B19']:assert w['Inputs'][cell].value is None
 assert 'not' in (MOD/'README.md').read_text().lower()
@pytest.mark.parametrize('f,k,t',[(-.01,.02,.25),(.05,.05,0),(.05,.03,1),(.01,-.01,.5)])
def test_put_call_parity_and_intrinsic(f,k,t):
 p=E.put(f,k,.008,t);call=E.put(-f,-k,.008,t)
 assert call-p==pytest.approx(f-k,abs=1e-12)
 assert p>=max(k-f,0)-1e-12
def test_delta_vega_independent_differences():
 h=1e-6;f=.05;k=.0475;sig=.008;t=.25
 delta=(E.put(f+h,k,sig,t)-E.put(f-h,k,sig,t))/(2*h)
 vega=(E.put(f,k,sig+h,t)-E.put(f,k,sig-h,t))/(2*h)
 d=(k-f)/(sig*math.sqrt(t))
 assert delta==pytest.approx(-E.cdf(d),abs=1e-8)
 assert vega==pytest.approx(math.sqrt(t)*E.pdf(d),abs=1e-8)
def test_normal_integration():
 n=20000;dx=16/n;f=.05;sig=.008;t=.25;width=.0025
 integrated=sum(min(max(-sig*math.sqrt(t)*(-8+(i+.5)*dx),0),width)*E.pdf(-8+(i+.5)*dx)*dx for i in range(n))
 assert integrated==pytest.approx(E.put(f,f,sig,t)-E.put(f,f-width,sig,t),abs=5e-10)
def test_bad_skew_flag_and_no_false_recommendation():
 r=E.calculate(short_vol_bp=400)
 assert not r['vertical_valid']
 assert 'ILLUSTRATIVE' in r['status']
@pytest.mark.parametrize('updates',[{'notional':0},{'tenor':2.5},{'expiry':0},{'width_bp':0},{'long_vol_bp':-1},{'vol_mode':1},{'forward_mode':1}])
def test_reject_invalid_or_missing(updates):
 with pytest.raises(ValueError):E.calculate(**updates)
def edit_xml(src,dst,values):
 ns={'m':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
 with zipfile.ZipFile(src) as z,zipfile.ZipFile(dst,'w',zipfile.ZIP_DEFLATED) as out:
  for obj in z.infolist():
   b=z.read(obj.filename)
   if obj.filename.startswith('xl/worksheets/sheet') and obj.filename.endswith('.xml'):
    root=ET.fromstring(b)
    for addr,v in (values.items() if obj.filename=='xl/worksheets/sheet2.xml' else []):
     c=root.find(f'.//m:c[@r="{addr}"]',ns)
     if c is None:raise AssertionError('Missing input cell '+addr)
     for child in list(c):c.remove(child)
     if v is None:c.attrib.pop('t',None)
     elif isinstance(v,str):
      c.set('t','inlineStr');st=ET.SubElement(c,'{'+ns['m']+'}is');ET.SubElement(st,'{'+ns['m']+'}t').text=v
     else:
      c.set('t','n');ET.SubElement(c,'{'+ns['m']+'}v').text=str(v)
    # File edits bypass Calc's dirty-cell tracking. Invalidate formula caches so
    # the native engine must recalculate every dependent output from the edits.
    for c in root.findall('.//m:c',ns):
     if c.find('m:f',ns) is not None:
      v=c.find('m:v',ns)
      if v is not None:c.remove(v)
    b=ET.tostring(root,encoding='utf-8',xml_declaration=True)
   out.writestr(obj,b)
def recalc(src,outdir):
 soffice=os.environ.get('SOFFICE')
 if not soffice:pytest.skip('Set SOFFICE for native LibreOffice change tests')
 profile=outdir/'profile'
 subprocess.run([soffice,f'-env:UserInstallation={profile.as_uri()}','--headless','--convert-to','xlsx','--outdir',str(outdir),str(src)],check=True,capture_output=True,timeout=90)
 return outdir/src.name
@pytest.mark.parametrize('cells,updates',[
 ({'B6':25000000,'B9':40,'B20':25,'B21':-50},{'notional':25000000,'width_bp':40,'curve_shock_bp':25,'scenario_shift_bp':-50}),
 ({'B7':.5,'B8':5,'B10':0,'B11':0},{'expiry':.5,'tenor':5,'long_vol_bp':0,'short_vol_bp':0}),
 ({'B11':100,'B12':2},{'short_vol_bp':100,'cost_bp':2}),
 ({'B13':1,'B14':.047,'B15':1,'B16':70,'B17':75,'B18':'2026-10-28 15:00 London - TEST','B19':'SYNTHETIC TEST INPUTS'},
  {'forward_mode':1,'dealer_forward':.047,'vol_mode':1,'dealer_long_vol_bp':70,'dealer_short_vol_bp':75,'quote_time':'test','quote_source':'test'})])
def test_native_input_changes(tmp_path,cells,updates):
 src=tmp_path/'changed.xlsx';edit_xml(BOOK,src,cells);out=tmp_path/'out';out.mkdir();parity(recalc(src,out),**updates)
def test_native_missing_quotes_block(tmp_path):
 src=tmp_path/'missing.xlsx';edit_xml(BOOK,src,{'B15':1});out=tmp_path/'out';out.mkdir()
 w=openpyxl.load_workbook(recalc(src,out),data_only=True)
 assert w['Summary']['B7'].value=='BLOCKED: INVALID / MISSING INPUTS'
 assert w['Pricing']['B24'].data_type=='e'
