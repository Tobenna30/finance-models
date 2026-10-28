from pathlib import Path
import importlib.util,sys,json,hashlib
import numpy as np
import openpyxl
import pytest
P=Path(__file__).resolve().parents[1]/'projects/03-systematic-quant/momentum'
def load():
 old=sys.modules.pop('assumptions',None);sys.path.insert(0,str(P));spec=importlib.util.spec_from_file_location('momentum_engine',P/'engine.py');e=importlib.util.module_from_spec(spec);spec.loader.exec_module(e);sys.path.pop(0);sys.modules.pop('assumptions',None)
 if old:sys.modules['assumptions']=old
 return e
E=load();R=E.results()
def test_source_hashes_and_first_month_value_weighted():
 for x in json.loads((P/'data/manifest.json').read_text()):assert hashlib.sha256((P/'data'/(x['name']+'.zip')).read_bytes()).hexdigest()==x['sha256']
 d=E.raw();assert d.index.is_unique;assert d.loc['1927-01','decile1']==pytest.approx(-.0332);assert d.loc['1927-01','decile10']==pytest.approx(-.0024);assert d.loc['1927-01','mom']==pytest.approx(.0057);assert d.index.max()=='2026-09'
def test_monthly_weights_costs_and_cash_tie():
 d=E.panel();assert np.allclose(d.gross,.5*(d.decile10-d.decile1));assert np.allclose(d.execution_cost,.001);assert np.allclose(d.borrow_cost,.01/12);assert np.allclose(d.net,d.gross-d.execution_cost-d.borrow_cost)
def test_cost_monotonicity_and_drag():
 low=E.stats(E.panel(0));base=E.stats(E.panel());high=E.stats(E.panel(25))
 assert low['net_sharpe']>base['net_sharpe']>high['net_sharpe'];assert base['annual_total_cost_drag']==pytest.approx(.022)
 assert base['annual_execution_drag']==pytest.approx(.012)
def test_drawdown_includes_starting_capital_and_recovery():
 d=E.panel();dd=E.drawdowns(d);assert dd[0]['trough']=='2021-05';assert dd[0]['recovery'] is None
 assert dd[0]['drawdown']==pytest.approx(d.drawdown.min());assert np.all(d.high_water>=1);assert d.wealth.iloc[-1]==pytest.approx(np.prod(1+d.net))
def test_chronological_evaluation_and_walk_forward_no_future():
 d=E.panel();wf=E.walk_forward(d);truncated=E.walk_forward(d.loc[:'2010-12'])
 for row in wf:
  assert row['train_end']<str(row['year']);assert row['train_months']>=444
 assert [x for x in wf if x['year']<=2010]==truncated
 assert R['design']['months']==444;assert R['evaluation']['months']==321
 assert R['design']['end']<R['evaluation']['start']
def test_2009_crash_and_published_factor_not_independent():
 d=E.panel();q=d.loc['2009-01':'2009-12'];assert R['crash_2009']['net_return']==pytest.approx(np.prod(1+q.net)-1);assert R['crash_2009']['net_return']<-.4
 assert .8<R['correlation_full']<1
def test_workbook_formula_cash_and_engine_parity():
 p=P.parent/'models/Momentum_Factor_Backtest.xlsx';w=openpyxl.load_workbook(p,data_only=True)
 for s in w:assert not [(c.coordinate,c.value) for row in s for c in row if c.data_type=='e']
 for i,key in [(2,'design'),(3,'evaluation'),(4,'full')]:
  for col,k in [('C','net_sharpe'),('F','net_cagr'),('G','annual_execution_drag'),('H','annual_borrow_drag'),('I','factor_correlation')]:assert w['Performance'][f'{col}{i}'].value==pytest.approx(R[key][k],abs=1e-10)
 assert w['Backtest']['G2'].value==pytest.approx(E.panel().net.iloc[0]);assert w['Performance']['J4'].value==pytest.approx(R['worst_drawdown']['drawdown'])
 f=openpyxl.load_workbook(p,data_only=False);assert f['Backtest']['G2'].data_type=='f';assert f['Backtest']['G766'].data_type=='f'
 for row in [2,6,10]:assert w['Sensitivity'].cell(row,4).value>w['Sensitivity'].cell(row+1,4).value
