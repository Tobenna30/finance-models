"""Conditional funding, claim allocation, Class B new-cash cost and Excel parity."""
import importlib
import json
import os
from pathlib import Path
import shutil
import subprocess
import tempfile
import pytest
from openpyxl import load_workbook
P='projects.07-restructuring.thames.'
e=importlib.import_module(P+'engine');a=importlib.import_module(P+'assumptions');b=importlib.import_module(P+'build_excel')

def test_facility_not_double_counted():
    r=e.calculate()
    assert sum(x['existing_draw'] for x in r['months'])==350
    assert sum(x['upsize_draw'] for x in r['months'])==2400
    assert r['new_draws']==2750 and r['ss_claim']==5400
    assert r['months'][-1]['month']=='2027-10-01'
    assert r['total_fees']==48

def test_cash_and_debt_rollforward():
    r=e.calculate();v=a.values()
    assert r['ending_cash']==pytest.approx(200+2750-12*200-48-r['total_interest'])
    assert r['ending_cash']==pytest.approx(342.546875)
    for i,x in enumerate(r['months']):
        assert x['closing_cash']==pytest.approx(x['opening_cash']+x['existing_draw']+x['upsize_draw']-200-x['fees']-x['incremental_interest']-x['principal'])
        if i:assert x['opening_cash']==r['months'][i-1]['closing_cash']
    # Opening super senior coupon is already inside cash use, not added twice.
    assert r['months'][0]['incremental_interest']==pytest.approx(750/2*.0975/12)

def test_approval_and_extension_are_separate():
    blocked=e.calculate({'funding_case':0});old=e.calculate({'funding_case':1});full=e.calculate();failed=e.calculate({'maturity_case':0})
    assert blocked['ending_cash'] < old['ending_cash'] < full['ending_cash']
    assert blocked['new_draws']==0 and old['new_draws']==350
    assert old['first_buffer_breach']=='2027-01-01'
    assert full['first_buffer_breach'] is None
    assert failed['ending_cash']==pytest.approx(full['ending_cash']-500)
    assert failed['a_claim']==16000  # maturity requirement is not proven redemption

def test_waterfall_conserves_value():
    for ev in (0,1000,18000,26000,30000):
        r=e.calculate({'ev':ev});w=r['waterfall']
        assert sum(w[k] for k in ('priority_paid','ss_paid','a_paid','b_paid','equity'))==pytest.approx(ev*.95)
        assert 0<=w['a_recovery']<=1 and 0<=w['b_recovery']<=1
    r=e.calculate();assert r['waterfall']['a_recovery']==pytest.approx(.6375)
    assert r['waterfall']['b_paid']==0
    par=e.calculate({'ev':r['a_break_even_ev']})
    assert par['waterfall']['a_recovery']==pytest.approx(1)
    assert par['waterfall']['b_paid']==pytest.approx(0,abs=1e-8)

def test_class_b_recovery_is_not_ownership_or_free_new_money():
    r=e.calculate()
    assert r['b_legacy_value']==59
    assert r['b_legacy_fraction']==pytest.approx(59/6000)
    assert r['b_legacy_fraction']!=.059
    assert r['b_new_fraction']==.025
    assert r['b_net_package']==59  # 59+150 asset value minus 150 paid
    down=e.calculate({'equity_stress':.5})
    assert down['b_marked_recovery']==pytest.approx(.0295)
    assert down['b_net_package']==pytest.approx(-45.5)
    assert r['b_full_par_equity']==pytest.approx(101694.915254)

def test_scenario_ordering_and_historical_reference():
    low=e.calculate({'burn':250});base=e.calculate();high=e.calculate({'burn':150})
    assert low['ending_cash']<base['ending_cash']<high['ending_cash']
    hist=e.calculate({'ev':base['historical_quote_implied_ev']})
    assert hist['waterfall']['a_recovery']==pytest.approx(.69)

def parity(w,r):
    for i,x in enumerate(r['months'],4):
        for row,key in {7:'opening_cash',8:'existing_draw',9:'upsize_draw',10:'fees',11:'incremental_interest',12:'principal',13:'closing_cash',14:'buffer_gap',16:'ss_debt'}.items():assert w['Liquidity'].cell(row,i).value==pytest.approx(x[key],abs=1e-7)
        assert w['Liquidity'].cell(19,i).value==pytest.approx(0,abs=1e-7)
    for row,key in {22:'minimum_cash',23:'peak_buffer_gap',24:'total_interest',25:'total_fees'}.items():assert w['Liquidity'][f'D{row}'].value==pytest.approx(r[key])
    for cell,key in {'D8':'distributable','E10':'priority_paid','E11':'ss_paid','E12':'a_paid','E13':'b_paid','D14':'equity','F12':'a_recovery','F13':'b_recovery'}.items():assert w['Recovery'][cell].value==pytest.approx(r['waterfall'][key],abs=1e-7)
    assert w['Recovery']['D16'].value==pytest.approx(0,abs=1e-7)
    assert w['Recovery']['D18'].value==pytest.approx(r['a_break_even_ev'])
    assert w['Recovery']['D19'].value==pytest.approx(r['historical_quote_implied_ev'])
    for row,key in {7:'b_legacy_value',8:'b_legacy_fraction',10:'b_new_fraction',12:'b_marked_legacy',13:'b_marked_new',14:'b_net_package',15:'b_marked_recovery',17:'b_full_par_equity'}.items():assert w['Class B'][f'D{row}'].value==pytest.approx(r[key],abs=1e-6)

def test_cached_workbook_and_live_formulas():
    w=load_workbook(b.OUT,data_only=True);parity(w,e.calculate())
    assert not [(s.title,c.coordinate) for s in w for row in s for c in row if c.data_type=='e']
    f=load_workbook(b.OUT,data_only=False)
    for c in range(4,16):assert all(f['Liquidity'].cell(row,c).data_type=='f' for row in (7,8,9,10,11,12,13,14,15,16,18,19))
    for key,v in a.INPUTS.items():assert f['Inputs'][f'D{b.I[key]}'].value==v[1] and v[2] in ('S','A','D') and v[4]
    assert json.loads((a.ROOT/'outputs/results.json').read_text())['upsize_assumed']['ending_cash']==pytest.approx(e.calculate()['ending_cash'])

@pytest.mark.parametrize('change',[{'funding_case':0},{'funding_case':1},{'maturity_case':0,'equity_stress':.5},{'ev':27000,'burn':250,'ss_rate':.12}])
def test_native_input_changes(change):
    soffice=os.environ.get('SOFFICE') or shutil.which('soffice')
    if not soffice:pytest.skip('LibreOffice unavailable; cached parity mandatory')
    with tempfile.TemporaryDirectory() as tmp:
        tmp=Path(tmp);out=tmp/'out';out.mkdir();w=load_workbook(b.OUT)
        for key,v in change.items():w['Inputs'][f'D{b.I[key]}']=v
        path=tmp/'change.xlsx';w.save(path)
        subprocess.run([soffice,'-env:UserInstallation='+(tmp/'profile').as_uri(),'--headless','--convert-to','xlsx:Calc MS Excel 2007 XML','--outdir',str(out),str(path)],capture_output=True,check=True,timeout=120)
        parity(load_workbook(out/path.name,data_only=True),e.calculate(change))

def test_invalid_inputs():
    for change in ({'funding_case':3},{'maturity_case':2},{'facility':2000},{'cost':1},{'equity':0},{'b_claim':0}):
        with pytest.raises(ValueError):e.calculate(change)
