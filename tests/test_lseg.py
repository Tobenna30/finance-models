import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "projects" / "01-equity-pitch-kit"))

from lseg.engine import run  # noqa: E402
from lseg import assumptions as A  # noqa: E402


def test_balance_sheet_balances_every_year_every_scenario():
    for s in (1, 2, 3):
        r = run(s)["r"]
        for y in A.YEARS_FCST:
            assert abs(r["check"][y]) < 1e-6, (s, y, r["check"][y])


def test_cash_flow_ties_to_balance_sheet():
    r = run(2)["r"]
    cash = A.OPENING_BS["Cash"]
    for y in A.YEARS_FCST:
        cash += r["eqfcf"][y] - r["div_cash"][y] - r["nci_div"][y] - r["bb"][y]
        assert abs(cash - r["cash"][y]) < 1e-6


def test_scenarios_are_ordered():
    fv = [run(s)["fair_value"] for s in (1, 2, 3)]
    assert fv[0] < fv[1] < fv[2]


def test_2026_consistent_with_company_guidance():
    r = run(2)["r"]
    assert r["eqfcf"][2026] >= 2780            # guide: at least £2.8bn
    margin_25 = A.HIST["Adj EBITDA"][2025] / 8986
    assert 0.0090 <= r["ebitda"][2026] / r["tot"][2026] - margin_25 <= 0.0150   # guide +100-120bp cc
    assert 1.5 <= r["lev"][2027] <= 2.5        # guide: upper end of 1.5-2.5x pro forma 2027


def test_excel_forecasts_are_formulas(tmp_path):
    from openpyxl import load_workbook
    import lseg.build_excel as b
    wb = load_workbook(b.OUT)
    m = wb["Model"]
    formulas = [c.value for row in m.iter_rows(min_row=6, min_col=5, max_col=9) for c in row if c.value is not None]
    assert formulas and all(isinstance(v, str) and v.startswith("=") for v in formulas)
