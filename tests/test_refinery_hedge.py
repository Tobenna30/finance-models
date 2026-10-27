import sys
from pathlib import Path
import pytest
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "projects" / "14-derivatives-hedging"))

from refinery import assumptions as A  # noqa: E402
from refinery import engine as E  # noqa: E402

XLSX = Path(__file__).resolve().parents[1] / "projects" / "14-derivatives-hedging" / "models" / "Refinery_Hedge.xlsx"


def test_market_data_ties_to_files():
    assert A.CRACK_DAYS[-1] == "2026-10-20"
    assert abs(A.SPOT["brent"] - 115.22) < 1e-9
    assert abs(A.SPOT["diesel_crack"] - (5.265 * 42 - 115.22)) < 1e-9
    assert abs(A.USD_RATE - 0.0388) < 1e-12
    assert all(v > 0 for v in A.NORMAL_VOL.values())


def test_sources_complete():
    assert all(u.startswith("https://") for _, u in A.SOURCES.values())


def test_put_call_parity_and_zero_cost():
    for model in ("bachelier", "black76"):
        for p, rows in E.strikes(model).items():
            for k in rows:
                assert abs(k["put_prem"] - k["call_prem"]) < 1e-6, (model, p)
                assert k["put"] < A.FORWARD[p] < k["call"]
    f, k, v, t, r = 40.0, 38.0, 44.7, 0.75, 0.0388
    import math
    for price in (E.bachelier, E.black76):
        vol = v if price is E.bachelier else v / f
        assert abs(price(f, k, vol, t, r) - price(f, k, vol, t, r, call=False) - math.exp(-r * t) * (f - k)) < 1e-9


def test_bachelier_collar_is_symmetric_and_handles_negative_cracks():
    for p, rows in E.strikes().items():
        for k in rows:
            assert abs((k["call"] - A.FORWARD[p]) - (A.FORWARD[p] - k["put"])) < 1e-6
    assert E.bachelier(-5, -6, 10, 1, 0.04) > 0
    with pytest.raises(ValueError):
        E.black76(-5, -6, 0.5, 1, 0.04)


def test_black76_skew_puts_call_further_out():
    kb = E.strikes("black76")["diesel"]
    assert all(k["call"] > A.FORWARD["diesel"] * (1 + A.PUT_DISTANCE) for k in kb)
    assert E.outage_table("black76")[0.8]["unbacked_cost"] < E.outage_table()[0.8]["unbacked_cost"]


def test_headline_numbers_match_closed_form():
    t = E.outage_table()
    for c in A.COVERS:
        assert abs(t[c]["fall30"] + (c * A.PUT_DISTANCE + (1 - c) * 0.30)) < 1e-9
        assert abs(t[c]["unbacked_cost"] - c * A.OUTAGE_DAYS / 365 * (A.OUTAGE_SPIKE - A.PUT_DISTANCE)) < 1e-9
    assert E.recommend() == 0.8
    assert abs(E.case(0, -0.3)["change"] + 0.30) < 1e-9


def test_scenarios_ordered():
    g = E.grid()
    for c in A.COVERS:
        vals = [g[c][s] for s in A.SHOCKS]
        assert all(a < b for a, b in zip(vals, vals[1:]))           # bear < base < bull
    for s in A.SHOCKS:
        if s < 0:
            assert all(g[a][s] < g[b][s] for a, b in zip(A.COVERS, A.COVERS[1:]))   # more cover, less loss


def test_outage_volumes():
    r = E.case(0.8, 0.3, outage=True)
    assert sum(x["lost"] for x in r["rows"]) == A.OUTAGE_DAYS
    assert [x["lost"] for x in r["rows"]][6:8] == [31, 11]


def test_fx_overlay_locks_rand_costs():
    fx = E.fx_overlay(0.8)
    assert abs(fx[0.0]["hedged"] - fx[0.0]["unhedged"]) < 1e-6
    assert fx[-0.15]["hedged"] > fx[-0.15]["unhedged"]               # rand strength hurts less when hedged
    assert fx[-0.15]["hedged"] < fx[0.0]["hedged"] < fx[0.15]["hedged"]


def test_monte_carlo_hedge_lifts_tail():
    mc = E.monte_carlo()
    assert mc["hedged_p05"] > mc["unhedged_p05"]


@pytest.fixture(scope="module")
def wb_values():
    from openpyxl import load_workbook
    return load_workbook(XLSX, data_only=True), load_workbook(XLSX)


def test_excel_ties_to_engine(wb_values):
    wv, _ = wb_values
    cover = {r[0]: r[1] for r in wv["Cover"].iter_rows(min_row=5, values_only=True) if r[0]}
    for k, v in E.summary().items():
        if k in cover:
            assert abs(cover[k] - v) < 1e-6 * max(1, abs(v)), k
    sel = E.case(0.8, 0.3, outage=True)
    assert abs(cover["Selected case: margin change"] - sel["change"]) < 1e-9
    assert abs(cover["Selected case: unbacked call cost"] - sel["unbacked_cost"]) < 1e-9
    assert abs(cover["Max abs Bachelier net premium ($/bbl)"]) < 1e-9
    assert abs(cover["Max abs Black-76 net premium ($/bbl)"]) < 1e-4
    g = E.grid()
    for i, c in enumerate(A.COVERS):
        for j, s in enumerate(A.SHOCKS):
            assert abs(wv["Scenarios"].cell(9 + i, 2 + j).value - g[c][s]) < 1e-9
    fx = E.fx_overlay(0.8)
    for k, sh in enumerate(A.FX_SHOCKS):
        assert abs(wv["FX"].cell(17, 7 + 2 * k).value - fx[sh]["unhedged"]) < 1e-6
        assert abs(wv["FX"].cell(17, 8 + 2 * k).value - fx[sh]["hedged"]) < 1e-6


def test_excel_model_cells_are_formulas(wb_values):
    _, wf = wb_values
    for sheet, rng in (("Pricing", "B5:I16"), ("Monthly", "C5:K16"), ("Scenarios", "B8:F12"), ("FX", "C5:L16")):
        vals = [c.value for row in wf[sheet][rng] for c in row if c.value is not None]
        assert vals and all(isinstance(v, str) and v.startswith("=") for v in vals), sheet
