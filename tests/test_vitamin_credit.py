import os
from pathlib import Path
import shutil
import subprocess
import sys

import pytest
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
P=ROOT/'projects'/'09-private-credit'
sys.path.insert(0,str(P))
from vitamin import assumptions as A, engine as E  # noqa: E402

XLSX=P/'models'/'Vitamin_Well_Credit.xlsx'


def test_sources_and_evidence_tags():
    assert all(url.startswith('https://') for _,url in A.SOURCES.values())
    assert all(x[2] in ('S','A','D') and x[4] for x in A.INPUTS.values())
    # Unverified pricing and policy thresholds must never be labelled sourced covenants.
    assert all(A.INPUTS[k][2]=='A' for k in ('spread','base_rate','max_leverage','min_cover','opening_cash','fx'))


@pytest.mark.parametrize('case',A.CASES)
def test_cash_and_debt_reconcile_without_hidden_funding(case):
    rows=E.run(case)
    for r in rows:
        assert abs(r['cash_check'])<1e-8
        assert abs(r['debt_check'])<1e-8
        assert 0<=r['revolver']<=A.RCF+1e-8
        assert r['debt']>=0
        assert not (r['draw']>0 and r['repay']>0)
        assert r['draw']<=A.RCF-r['opening_rcf']+1e-8
        assert r['gross']==pytest.approx(r['debt']+r['revolver'])
        assert r['liquidity']==pytest.approx(r['cash']+A.RCF-r['revolver'])
    assert rows[0]['transaction']==824
    assert all(r['transaction']==0 for r in rows[1:])


def test_case_order_and_downside_policy_failure():
    d,b,u=(E.run(c) for c in ('Downside','Base','Bull'))
    for rd,rb,ru in zip(d,b,u):
        assert rd['cash_ebitda']<rb['cash_ebitda']<ru['cash_ebitda']
        assert rd['leverage']>rb['leverage']>ru['leverage']
        assert rd['cover']<rb['cover']<ru['cover']
    assert all(r['lev_headroom']>=0 and r['cover_headroom']>=0 for r in b)
    assert min(r['lev_headroom'] for r in d)<0
    assert min(r['cover_headroom'] for r in d)<0
    assert max(r['unfunded'] for r in d)==0


def test_capacity_binds_and_ebitda_boundary():
    c=E.capacity()
    assert c['capacity']==min(c['leverage'],c['cover'])
    assert c['excess']<0
    assert E.capacity(c['break_even_ebitda'])['leverage']==pytest.approx(A.TLB)
    assert E.capacity(c['break_even_ebitda']*.99)['excess']>0
    assert E.capacity(rate=.15)['capacity']<c['capacity']
    with pytest.raises(ValueError):E.capacity(rate=0)


def test_recovery_waterfall_priority_and_bounds():
    for m in (0,4,6,8,30):
        r=E.recovery(multiple=m)
        assert 0<=r['recovery']<=1
        assert r['tlb_proceeds']+r['rcf_proceeds']+r['equity']==pytest.approx(r['pool'])
        assert r['shortfall']>=0
    assert E.recovery(multiple=0)['recovery']==0
    assert E.recovery(multiple=30)['recovery']==1
    assert E.recovery(rcf_draw=1)['recovery']<E.recovery(rcf_draw=0)['recovery']
    assert E.recovery(haircut=.5)['recovery']<E.recovery(haircut=.2)['recovery']
    with pytest.raises(ValueError):E.recovery(haircut=1.1)


@pytest.fixture(scope='module')
def wb():
    return load_workbook(XLSX,data_only=True),load_workbook(XLSX)


MAP={
  'Revenue': 'revenue','Adjusted EBITDA':'ebitda','Cash EBITDA':'cash_ebitda',
  'Interest on opening balances':'interest','Undrawn commitment fee':'fee','D&A proxy':'da',
  'Cash tax proxy':'tax','Capex':'capex','Working capital investment':'nwc','Transaction costs':'transaction',
  'FCF after interest and tax':'fcf','Mandatory term-loan amortisation':'amort',
  'Revolver draw':'draw','Revolver repayment':'repay','Term-loan cash sweep':'sweep',
  'Closing term loan':'debt','Closing revolver':'revolver','Closing cash':'cash',
  'Gross funded debt':'gross','Net funded debt':'net','Cash plus undrawn revolver':'liquidity',
  'Gross leverage / cash EBITDA':'leverage','Cash interest cover':'cover',
  'Headroom to assumed leverage ceiling':'lev_headroom','Headroom above assumed cover floor':'cover_headroom',
  'Unfunded cash-minimum gap':'unfunded','Cash rollforward check':'cash_check','Debt rollforward check':'debt_check'}

def assert_forecast(w,case):
    ws=w['Forecast'];labels={ws.cell(r,1).value:r for r in range(5,39)}
    for lab,key in MAP.items():
        for j,row in enumerate(E.run(case),2):
            assert ws.cell(labels[lab],j).value==pytest.approx(row[key],abs=1e-7), (case,lab,j)


def test_excel_engine_parity_and_forecast_formulas(wb):
    values,forms=wb
    assert_forecast(values,'Base')
    for row in forms['Forecast'].iter_rows(min_row=5,max_row=38,min_col=2,max_col=6):
        assert all(isinstance(c.value,str) and c.value.startswith('=') for c in row)
    cap=E.capacity();rec=E.recovery()
    assert values['Capacity']['B13'].value==pytest.approx(cap['capacity'])
    assert values['Capacity']['B16'].value==pytest.approx(1-cap['break_even_ebitda']/A.EBITDA)
    assert values['Recovery']['B13'].value==pytest.approx(rec['recovery'])
    assert values['Sensitivity']['D7'].value==pytest.approx(rec['recovery'])
    assert values['Sensitivity']['C16'].value==pytest.approx(cap['capacity'])


def test_no_cached_formula_errors(wb):
    values,_=wb
    errors=[(ws.title,c.coordinate,c.value) for ws in values for row in ws for c in row if c.data_type=='e']
    assert errors==[]


@pytest.mark.parametrize('case_number,case',[(1,'Downside'),(3,'Bull')])
def test_case_selector_recalculates_full_model(tmp_path,case_number,case):
    soffice=os.environ.get('SOFFICE') or shutil.which('soffice')
    bundled=Path.home()/'.cache/codex-runtimes/codex-primary-runtime/dependencies/bin/override/soffice'
    soffice=soffice or (str(bundled) if bundled.exists() else None)
    if not soffice:pytest.skip('LibreOffice required for native case-switch verification; engine/cached parity still tested')
    w=load_workbook(XLSX);w['Inputs']['C4']=case_number
    incoming=tmp_path/'incoming';out=tmp_path/'recalc';incoming.mkdir();out.mkdir()
    file=incoming/'case.xlsx';w.save(file)
    subprocess.run([soffice,'--headless','--convert-to','xlsx:Calc MS Excel 2007 XML','--outdir',str(out),str(file)],
                   check=True,capture_output=True,timeout=120)
    recalc=load_workbook(out/'case.xlsx',data_only=True)
    assert recalc['Forecast']['B3'].value==case
    assert_forecast(recalc,case)
    # Static source inputs retain their meaning across the selector.
    assert recalc['Capacity']['B5'].value==A.TLB


def test_exhausted_revolver_leaves_visible_unfunded_gap(monkeypatch):
    old=A.INPUTS['transaction_cost']
    monkeypatch.setitem(A.INPUTS,'transaction_cost',(old[0],10000.0,*old[2:]))
    row=E.run('Downside')[0]
    assert row['revolver']==pytest.approx(A.RCF)
    assert row['unfunded']>0
    assert row['cash']<A.value('minimum_cash')
    assert abs(row['cash_check'])<1e-8
