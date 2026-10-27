"""Build the live Darktrace workbook (04 sell-side + 08 LBO). Every forecast and valuation cell is a formula
that points back to the Inputs sheet. Run from projects/04-08-darktrace: python -m darktrace.build_excel
Then recalculate with LibreOffice: python -m darktrace.recalc"""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter as CL
from openpyxl.chart import BarChart, Reference
from . import assumptions as A

BOLD = Font(bold=True)
HDR, HDRF = PatternFill("solid", fgColor="1F3864"), Font(bold=True, color="FFFFFF")
SEC = PatternFill("solid", fgColor="D9E1F2")
M, M1, PCT, X, P0, USD = '#,##0;(#,##0);"-"', '#,##0.0;(#,##0.0);"-"', '0.0%;(0.0%);"-"', '0.0"x"', '#,##0"p"', '$0.00'
RAW = Path(__file__).resolve().parents[1] / "models" / "_raw_Darktrace_04_08.xlsx"
OUT = Path(__file__).resolve().parents[1] / "models" / "Darktrace_04_08.xlsx"

wb = Workbook()
cover = wb.active; cover.title = "Cover"
names = ["Inputs", "History", "Offer", "DCF", "Comps", "Precedents", "Football field", "Rival bid",
         "LBO", "Downside", "LBO grid", "Ability to pay", "Sources"]
S = {n: wb.create_sheet(n) for n in names}


def title(ws, t, sub=None):
    ws["A1"] = t; ws["A1"].font = Font(bold=True, size=14)
    if sub: ws["A2"] = sub


def hdr(ws, r, c, v):
    x = ws.cell(r, c, v); x.fill = HDR; x.font = HDRF; return x


def put(ws, ref, v, fmt=None, bold=False):
    c = ws[ref]; c.value = v
    if isinstance(v, str) and v.startswith("="):
        c.font = Font(color="008000", bold=bold) if "!" in v else Font(bold=bold)
    elif isinstance(v, (int, float)):
        c.font = Font(color="0000FF", bold=bold)
    elif bold:
        c.font = BOLD
    if fmt: c.number_format = fmt
    return c


def src_label(k):
    return f"{k}: {A.SOURCES[k][0]}" if k in A.SOURCES else k


# =============================================================== Inputs
inp = S["Inputs"]; I = {}
title(inp, "Inputs", "Blue = hard-coded input. Tag: S = sourced, A = assumption, D = derived. $m unless stated. "
      "Series columns F onwards = FY25 onwards.")
for j, h in enumerate(["Input", "Unit", "Value", "Tag", "Source / note"], start=1): hdr(inp, 4, j, h)
for j, y in enumerate(A.YEARS_DCF): hdr(inp, 4, 6 + j, f"FY{y % 100}")
row = 5


def section(t):
    global row
    row += 1; inp[f"A{row}"] = t; inp[f"A{row}"].font = BOLD
    for c in range(1, 13): inp.cell(row, c).fill = SEC
    row += 1


def scalar(key, label, val, unit, tag, note="", fmt=None):
    global row
    inp[f"A{row}"], inp[f"B{row}"], inp[f"D{row}"], inp[f"E{row}"] = label, unit, tag, note
    put(inp, f"C{row}", val, fmt); I[key] = f"Inputs!$C${row}"; row += 1


def series(key, label, vals, unit, tag, note="", fmt=PCT):
    global row
    inp[f"A{row}"], inp[f"B{row}"], inp[f"D{row}"], inp[f"E{row}"] = label, unit, tag, note
    I[key] = []
    for j, v in enumerate(vals):
        ref = f"{CL(6 + j)}{row}"; put(inp, ref, v, fmt); I[key].append(f"Inputs!${CL(6 + j)}${row}")
    row += 1


D = A.DEAL
section("Offer terms (Rule 2.7 announcement, 26 Apr 2024)")
scalar("offer_usd", "Offer price per share", D["offer_usd"], "$", "S", src_label("R27"), USD)
scalar("offer_p", "Offer price, GBP equivalent at announcement", D["offer_p"], "p", "S", src_label("R27"), P0)
scalar("eqv", "Equity value, fully diluted", D["equity_value"], "$m", "S", src_label("R27"), M)
scalar("ev", "Enterprise value", D["enterprise_value"], "$m", "S", src_label("R27"), M)
scalar("ltm_rev", "LTM revenue (12m to 31 Dec 2023)", D["ltm_revenue"], "$m", "S", src_label("R27"), M)
scalar("ltm_ebitda", "LTM adjusted EBITDA (12m to 31 Dec 2023)", D["ltm_ebitda"], "$m", "S", src_label("R27"), M)
scalar("ev_rev_stated", "Stated EV / LTM revenue", D["ev_rev_stated"], "x", "S", src_label("R27"), X)
scalar("ev_ebitda_stated", "Stated EV / LTM adjusted EBITDA", D["ev_ebitda_stated"], "x", "S", src_label("R27"), X)
scalar("undisturbed", "Undisturbed close, 25 Apr 2024", D["undisturbed_p"], "p", "S", src_label("R27"), P0)
scalar("vwap3m", "3-month VWAP to 25 Apr 2024", D["vwap3m_p"], "p", "S", src_label("R27"), '0.0"p"')
scalar("high12m", "12-month high close", D["high12m_p"], "p", "S", src_label("R27"), '0.0"p"')
scalar("placing", "Secondary placing price, 21 Mar 2024", D["placing_p"], "p", "S", src_label("R27"), P0)
scalar("ipo", "IPO price, 30 Apr 2021", D["ipo_p"], "p", "S", src_label("R27"), P0)
scalar("prem_close_s", "Stated premium to undisturbed close", D["prem_close"], "%", "S", src_label("R27"), PCT)
scalar("prem_vwap_s", "Stated premium to 3-month VWAP", D["prem_vwap3m"], "%", "S", src_label("R27"), PCT)
scalar("irrev", "Irrevocable undertakings (binding if a higher offer is made)", D["irrevocables_pct"], "% ISC", "S", src_label("SCH"), PCT)
scalar("court_for", "Court Meeting: Scheme Shares voted in favour", D["court_for_pct"], "%", "S", src_label("VOTE"), '0.00%')
scalar("court_turnout", "Court Meeting: Scheme Shares voted / issued capital", D["court_turnout_pct"], "%", "S", src_label("VOTE"), '0.00%')
scalar("cash_dec23", "Cash and equivalents, 31 Dec 2023", A.H1["cash_dec23"], "$m", "S", src_label("H124"), M1)
scalar("leases_dec23", "Lease liabilities, 31 Dec 2023", A.H1["leases_dec23"], "$m", "S", src_label("H124"), M1)

section("History (FY to 30 June)")
scalar("rev22", "Revenue FY22 (FY23 less reported $129.9m increase)", A.HIST["Revenue"][2022], "$m", "D", src_label("FY23"), M1)
scalar("rev23", "Revenue FY23", A.HIST["Revenue"][2023], "$m", "S", src_label("FY23"), M1)
scalar("rev24", "Revenue FY24 ('at least', unaudited)", A.HIST["Revenue"][2024], "$m", "S", src_label("Q424"), M1)
scalar("arr23", "ARR FY23, constant currency", A.HIST["ARR (cc)"][2023], "$m", "S", src_label("FY23"), M1)
scalar("arr24", "ARR FY24, constant currency", A.HIST["ARR (cc)"][2024], "$m", "S", src_label("Q424"), M1)
scalar("nrr24", "Net ARR retention FY24", A.HIST["Net ARR retention"][2024], "%", "S", src_label("Q424"), PCT)
scalar("churn24", "Gross ARR churn FY24", A.HIST["Gross churn"][2024], "%", "S", src_label("Q424"), PCT)
scalar("rev_h1_24", "Revenue H1 FY24", A.H1["rev_h1fy24"], "$m", "S", src_label("H124"), M1)
scalar("rev_h1_23", "Revenue H1 FY23", A.H1["rev_h1fy23"], "$m", "S", src_label("H124"), M1)
scalar("ebitda_h1_24", "Adjusted EBITDA H1 FY24", A.H1["ebitda_h1fy24"], "$m", "S", src_label("H124"), M1)
scalar("ebitda_h1_23", "Adjusted EBITDA H1 FY23 (restated definition)", A.H1["ebitda_h1fy23"], "$m", "S", src_label("H124"), M1)
scalar("adj_ebit_h1_24", "Adjusted EBIT H1 FY24", A.H1["adj_ebit_h1fy24"], "$m", "S", src_label("H124"), M1)
scalar("sbp_h1_24", "Share-based payment charge H1 FY24", A.H1["sbp_h1fy24"], "$m", "S", src_label("H124"), M1)
scalar("capex_h1_24", "Capex H1 FY24 (PPE $7.0m + capitalised development $1.0m)", A.H1["capex_h1fy24"], "$m", "S", src_label("H124"), M1)
scalar("ebitda_m24", "FY24 adjusted EBITDA margin", A.FY24_EBITDA_MARGIN, "%", "A", "Guide was at least 21% (H124); H1 delivered 25.6%", PCT)
scalar("tax", "UK corporation tax rate", A.TAX, "%", "S", src_label("UKTAX"), PCT)
scalar("cir", "Corporate interest restriction: cap as % of tax-EBITDA", A.CIR_CAP, "%", "S", src_label("UKTAX"), PCT)
scalar("nwc", "Working capital inflow, % of revenue increase", A.DCF["nwc_pct_drev"], "%", "A", "Deferred revenue builds with billings", PCT)

section("Project 04: standalone DCF (valuation at 1 Jul 2024, mid-year discounting)")
series("g_dcf", "Revenue growth", A.DCF["rev_growth"], "%", "A", "FY24 +26%, ARR +22.7%; net ARR added flat y/y, so growth fades")
series("m_dcf", "Adjusted EBITDA margin", A.DCF["ebitda_margin"], "%", "A", "Steady-state adj EBIT target 'mid-20%s' (FY23) plus D&A of c.4%")
series("sbp_dcf", "Share-based payments, % revenue (treated as a cost)", A.DCF["sbp_pct"], "%", "A", "H1 FY24 6.9% of revenue")
scalar("rf", "Risk-free rate: US 10-year, Apr 2024 average", A.DCF["rf"], "%", "S", src_label("UST"), '0.00%')
scalar("beta", "Equity beta", A.DCF["beta"], "x", "A", "No debt; listed cyber peers trade above market beta", '0.00')
scalar("erp", "Equity risk premium", A.DCF["erp"], "%", "A", "", PCT)
scalar("tvg", "Terminal growth", A.DCF["tv_growth"], "%", "A", "", PCT)
scalar("sbp_switch", "Deduct SBP in free cash flow (1 = yes, 0 = no)", 1, "1/0", "A", "Set 0 to see the value with SBP added back", "0")

section("Project 04: rival bid screen")
scalar("syn_pct", "Strategic cost synergies, % FY25E revenue", A.RIVAL["synergy_pct_rev"], "%", "A", "Sales and G&A overlap", PCT)
scalar("syn_mult", "Multiple on after-tax synergies", A.RIVAL["synergy_multiple"], "x", "A", "", X)
scalar("syn_cost", "One-off integration cost, x run-rate synergies", A.RIVAL["integration_cost_x"], "x", "A", "", X)
scalar("syn_share", "Share of synergy value paid away to win", A.RIVAL["strategic_share_kept"], "%", "A", "", PCT)
scalar("pe_irr", "Rival sponsor IRR hurdle", A.RIVAL["competing_pe_irr"], "%", "A", "", PCT)

section("Project 08: debt financing at close")
scalar("tlb", "First-lien term loan", A.DEBT["tlb"], "$m", "S", src_label("PB") + " (interim facility at announcement: $1,685m)", M)
scalar("tlb_m", "First-lien margin over SOFR", A.DEBT["tlb_margin"], "%", "S", src_label("PB"), '0.00%')
scalar("tlb_oid", "First-lien OID (price)", A.DEBT["tlb_oid"], "price", "S", src_label("PB"), '0.0%')
scalar("tlb_tenor", "First-lien tenor", A.DEBT["tlb_tenor"], "years", "S", src_label("PB"), "0")
scalar("tlb_am", "First-lien amortisation, % of original face p.a.", A.DEBT["tlb_amort"], "%", "A", "Market norm for a US term loan B", PCT)
scalar("sl", "Second-lien term loan", A.DEBT["sl"], "$m", "S", src_label("PB") + " (interim: $460m)", M)
scalar("sl_m", "Second-lien margin over SOFR", A.DEBT["sl_margin"], "%", "S", src_label("PB"), '0.00%')
scalar("sl_oid", "Second-lien OID (price)", A.DEBT["sl_oid"], "price", "S", src_label("PB"), '0.0%')
scalar("sl_tenor", "Second-lien tenor", A.DEBT["sl_tenor"], "years", "S", src_label("PB"), "0")
series("sofr", "SOFR path FY25-FY29", A.SOFR, "%", "A", "From c.4.85% at close, easing", '0.00%')
scalar("fin_fee", "Financing fees, % of debt", A.FEES["financing_pct"], "%", "A", "", PCT)
scalar("adv_fee", "Advisory and other deal fees, % of EV", A.FEES["advisory_pct_ev"], "%", "A", "", PCT)
scalar("cash_close", "Target cash at close (31 Dec 2023 balance as proxy)", A.CASH_AT_CLOSE, "$m", "A", "Last published balance", M1)
scalar("min_cash", "Minimum operating cash", A.MIN_CASH, "$m", "A", "", M)
scalar("sweep", "Excess cash swept to first-lien", A.SWEEP, "%", "A", "", PCT)
scalar("lease", "Annual lease cash payments", A.LEASE_CASH, "$m", "A", "Leases $60m at Dec 2023", M)
scalar("mip", "Management incentive plan, % of exit equity", A.MIP, "%", "A", "", PCT)
scalar("leases_exit", "Lease liabilities deducted at exit (held flat)", A.LEASES_EXIT, "$m", "A", "Same treatment as the entry EV bridge", M1)
scalar("stub", "FY25 months under sponsor ownership / 12", A.CLOSE_STUB, "x", "D", "Effective 1 Oct 2024 (EFF)", "0.00")
scalar("hold", "Holding period to 30 Jun 2029", A.HOLD_YEARS, "years", "D", "", "0.00")
for case, ck in (("Base", "b"), ("Downside", "d")):
    section(f"Project 08: {case} operating case (FY25-FY29)")
    series(f"g_{ck}", "Revenue growth", A.CASES[case]["rev_growth"], "%", "A", "")
    series(f"m_{ck}", "Adjusted EBITDA margin", A.CASES[case]["ebitda_margin"], "%", "A",
           "Sponsor cost programme" if case == "Base" else "Limited cost take-out, slower growth")
    scalar(f"x_{ck}", "Exit EV / LTM EBITDA", A.CASES[case]["exit_multiple"], "x", "A", "", X)
section("Derived")
scalar("fx", "Implied USD per GBP ($7.75 / 620p)", "=" + I["offer_usd"] + "/(" + I["offer_p"] + "/100)", "$/GBP", "D", "", "0.0000")
scalar("fd", "Fully diluted shares (equity value / offer price)", "=" + I["eqv"] + "/" + I["offer_usd"], "m", "D", "", M1)
scalar("netcash", "Net cash in EV bridge (equity value less EV)", "=" + I["eqv"] + "-" + I["ev"], "$m", "D",
       "Equals cash $383.2m less leases $60.2m at 31 Dec 2023", M1)
scalar("capex_pct", "Capex, % revenue (H1 FY24)", "=" + I["capex_h1_24"] + "/" + I["rev_h1_24"], "%", "D", "", PCT)
scalar("da_pct", "D&A below adjusted EBITDA, % revenue (H1 FY24)",
       f"=({I['ebitda_h1_24']}-{I['adj_ebit_h1_24']})/{I['rev_h1_24']}", "%", "D", "", PCT)
scalar("wacc", "Cost of equity = WACC (no debt)", f"={I['rf']}+{I['beta']}*{I['erp']}", "%", "D", "", '0.00%')
scalar("ntm_w", "NTM weight on FY24 at 25 Apr 2024", A.NTM_WEIGHT_FY24, "x", "D", "May-Jun 2024 = 2 of 12 months", "0.000")
inp.column_dimensions["A"].width = 58; inp.column_dimensions["E"].width = 70
for c in "BCD": inp.column_dimensions[c].width = 10


def pps(ev_ref):
    """pence per share from an EV formula string."""
    return f"(({ev_ref})+{I['netcash']})/{I['fd']}*100/{I['fx']}"


# =============================================================== History
h = S["History"]
title(h, "Darktrace history (FY to 30 June) and LTM bridge")
for j, t in enumerate(["$m", "FY22", "FY23", "FY24"], start=1): hdr(h, 4, j, t)
put(h, "A5", "Revenue"); put(h, "B5", f"={I['rev22']}", M1); put(h, "C5", f"={I['rev23']}", M1); put(h, "D5", f"={I['rev24']}", M1)
put(h, "A6", "Growth"); put(h, "C6", "=C5/B5-1", PCT); put(h, "D6", "=D5/C5-1", PCT)
put(h, "A7", "ARR (constant currency)"); put(h, "C7", f"={I['arr23']}", M1); put(h, "D7", f"={I['arr24']}", M1)
put(h, "A8", "ARR growth"); put(h, "D8", "=D7/C7-1", PCT)
put(h, "A9", "Adjusted EBITDA (FY23 derived on the new definition; FY24 assumed margin)")
put(h, "C9", f"={I['ltm_ebitda']}-{I['ebitda_h1_24']}+{I['ebitda_h1_23']}", M1); put(h, "D9", f"=D5*{I['ebitda_m24']}", M1)
put(h, "A10", "Margin"); put(h, "C10", "=C9/C5", PCT); put(h, "D10", "=D9/D5", PCT)
put(h, "A12", "LTM to 31 Dec 2023", bold=True)
put(h, "A13", "Revenue: FY23 + H1 FY24 - H1 FY23"); put(h, "B13", f"={I['rev23']}+{I['rev_h1_24']}-{I['rev_h1_23']}", M1)
put(h, "A14", "Stated in the Rule 2.7 announcement"); put(h, "B14", f"={I['ltm_rev']}", M1)
put(h, "A15", "SBP as % revenue, H1 FY24"); put(h, "B15", f"={I['sbp_h1_24']}/{I['rev_h1_24']}", PCT)
put(h, "A16", "Adjusted EBITDA margin, H1 FY24"); put(h, "B16", f"={I['ebitda_h1_24']}/{I['rev_h1_24']}", PCT)
put(h, "A17", "Net ARR retention FY24 / gross churn FY24"); put(h, "B17", f"={I['nrr24']}", PCT); put(h, "C17", f"={I['churn24']}", PCT)
h.column_dimensions["A"].width = 64

# =============================================================== Offer
o = S["Offer"]
title(o, "Offer analysis: what Thoma Bravo paid", "All figures recomputed from inputs and checked against the stated numbers")
rows = [
    ("Offer price ($)", f"={I['offer_usd']}", USD), ("Offer price (p)", f"={I['offer_p']}", P0),
    ("Implied FX ($ per GBP)", f"={I['fx']}", "0.0000"), ("Fully diluted shares (m)", f"={I['fd']}", M1),
    ("Equity value ($m)", f"={I['eqv']}", M), ("Net cash ($m)", f"={I['netcash']}", M1),
    ("Enterprise value ($m)", f"={I['ev']}", M),
    ("EV / LTM revenue", f"={I['ev']}/{I['ltm_rev']}", X), ("Stated EV / LTM revenue", f"={I['ev_rev_stated']}", X),
    ("EV / LTM adjusted EBITDA", f"={I['ev']}/{I['ltm_ebitda']}", X), ("Stated EV / LTM EBITDA", f"={I['ev_ebitda_stated']}", X),
    ("EV / FY24 ARR", f"={I['ev']}/{I['arr24']}", X),
    ("Premium to undisturbed close (517p)", f"={I['offer_p']}/{I['undisturbed']}-1", PCT),
    ("Stated premium to close", f"={I['prem_close_s']}", PCT),
    ("Premium to 3-month VWAP (429.9p)", f"={I['offer_p']}/{I['vwap3m']}-1", PCT), ("Stated premium to VWAP", f"={I['prem_vwap_s']}", PCT),
    ("Premium to 12-month high (518.6p)", f"={I['offer_p']}/{I['high12m']}-1", PCT),
    ("Premium to March 2024 placing (425p)", f"={I['offer_p']}/{I['placing']}-1", PCT),
    ("Multiple of IPO price (250p)", f"={I['offer_p']}/{I['ipo']}", X),
    ("Run-up: undisturbed close vs 3-month VWAP", f"={I['undisturbed']}/{I['vwap3m']}-1", PCT),
    ("Irrevocables (% of issued capital)", f"={I['irrev']}", PCT),
    ("Court Meeting: % of Scheme Shares voted in favour", f"={I['court_for']}", '0.00%'),
    ("Court Meeting: turnout (% of issued capital)", f"={I['court_turnout']}", '0.00%'),
]
OFR = {}
for i, (lab, f, fmt) in enumerate(rows, start=4):
    o[f"A{i}"] = lab; put(o, f"B{i}", f, fmt); OFR[lab] = f"Offer!$B${i}"
r0 = 4 + len(rows) + 1
o[f"A{r0}"] = "Timetable"; o[f"A{r0}"].font = BOLD
for i, (lab, v) in enumerate([("Announcement (Rule 2.7)", D["announce_date"]), ("Scheme Document", D["scheme_doc_date"]),
                              ("Court Meeting and General Meeting", D["meeting_date"]), ("Regulatory conditions satisfied", D["regulatory_cleared"]),
                              ("Court sanction", D["sanction_date"]), ("Effective; delisted next day", D["effective_date"]),
                              ("Long stop date", D["long_stop"]), ("Rule 3 advisers to the board", D["rule3_advisers"]),
                              ("Bidco financial adviser", D["bidco_adviser"])], start=r0 + 1):
    o[f"A{i}"], o[f"B{i}"] = lab, v
o.column_dimensions["A"].width = 52; o.column_dimensions["B"].width = 30

# =============================================================== DCF
d = S["DCF"]; YD = A.YEARS_DCF; C0 = 3
title(d, "Standalone DCF (public-company case, before any bid)", "$m. Mid-year discounting from 1 Jul 2024. SBP is treated as a cost.")
hdr(d, 4, 1, "$m"); hdr(d, 4, 2, "FY24A")
for j, y in enumerate(YD): hdr(d, 4, C0 + j, f"FY{y % 100}E")
DR = {k: r for r, k in enumerate(["rev", "g", "ebitda", "m", "sbp", "da", "ebit", "tax", "capex", "dnwc", "ufcf", "t", "df", "pv"], start=5)}
labels = {"rev": "Revenue", "g": "Growth", "ebitda": "Adjusted EBITDA", "m": "Margin", "sbp": "Share-based payments",
          "da": "D&A", "ebit": "EBIT after SBP", "tax": "Tax", "capex": "Capex", "dnwc": "Working capital inflow",
          "ufcf": "Unlevered free cash flow", "t": "Discount period (mid-year)", "df": "Discount factor", "pv": "PV of UFCF"}
for k, r in DR.items(): d[f"A{r}"] = labels[k]
put(d, f"B{DR['rev']}", "=History!D5", M1)
for j, y in enumerate(YD):
    c = CL(C0 + j); p = CL(C0 + j - 1)
    put(d, f"{c}{DR['g']}", f"={I['g_dcf'][j]}", PCT)
    put(d, f"{c}{DR['rev']}", f"={p}{DR['rev']}*(1+{c}{DR['g']})", M1)
    put(d, f"{c}{DR['m']}", f"={I['m_dcf'][j]}", PCT)
    put(d, f"{c}{DR['ebitda']}", f"={c}{DR['rev']}*{c}{DR['m']}", M1)
    put(d, f"{c}{DR['sbp']}", f"={c}{DR['rev']}*{I['sbp_dcf'][j]}", M1)
    put(d, f"{c}{DR['da']}", f"={c}{DR['rev']}*{I['da_pct']}", M1)
    put(d, f"{c}{DR['ebit']}", f"={c}{DR['ebitda']}-{c}{DR['sbp']}-{c}{DR['da']}", M1)
    put(d, f"{c}{DR['tax']}", f"=MAX(0,{c}{DR['ebit']})*{I['tax']}", M1)
    put(d, f"{c}{DR['capex']}", f"={c}{DR['rev']}*{I['capex_pct']}", M1)
    put(d, f"{c}{DR['dnwc']}", f"=({c}{DR['rev']}-{p}{DR['rev']})*{I['nwc']}", M1)
    put(d, f"{c}{DR['ufcf']}", f"={c}{DR['ebitda']}-{c}{DR['sbp']}*{I['sbp_switch']}-{c}{DR['tax']}-{c}{DR['capex']}+{c}{DR['dnwc']}", M1, True)
    put(d, f"{c}{DR['t']}", f"={j}+0.5", "0.0")
    put(d, f"{c}{DR['df']}", f"=1/(1+{I['wacc']})^{c}{DR['t']}", "0.000")
    put(d, f"{c}{DR['pv']}", f"={c}{DR['ufcf']}*{c}{DR['df']}", M1)
LC = CL(C0 + len(YD) - 1); FC = CL(C0); N = len(YD)
val = [("WACC", f"={I['wacc']}", '0.00%'), ("Terminal growth", f"={I['tvg']}", PCT),
       ("Sum of PV of UFCF", f"=SUM({FC}{DR['pv']}:{LC}{DR['pv']})", M1),
       ("Terminal value (Gordon growth)", f"={LC}{DR['ufcf']}*(1+B22)/(B21-B22)", M1),
       ("PV of terminal value", f"=B24/(1+B21)^{N}", M1), ("Enterprise value", "=B23+B25", M1),
       ("Net cash", f"={I['netcash']}", M1), ("Equity value", "=B26+B27", M1),
       ("Value per share ($)", f"=B28/{I['fd']}", USD), ("Value per share (p)", f"=B29*100/{I['fx']}", P0),
       ("vs offer 620p", f"=B30/{I['offer_p']}-1", PCT), ("vs undisturbed 517p", f"=B30/{I['undisturbed']}-1", PCT),
       ("Terminal value as % of EV", "=B25/B26", PCT), ("Implied terminal EV / EBITDA", f"=B24/{LC}{DR['ebitda']}", X),
       ("Offer EV / FY25E EBITDA", f"={I['ev']}/{FC}{DR['ebitda']}", X)]
for i, (lab, f, fmt) in enumerate(val, start=21):
    d[f"A{i}"] = lab; put(d, f"B{i}", f, fmt, lab.startswith("Value per share (p)"))
DCFP = "DCF!$B$30"
g0 = 38; d[f"A{g0}"] = "Value per share (p): WACC (rows) x terminal growth (columns)"; d[f"A{g0}"].font = BOLD
for j, ga in enumerate(A.DCF_GRID_G):
    put(d, f"{CL(2 + j)}{g0 + 1}", f"={I['tvg']}+{ga}", PCT)
for i, wa in enumerate(A.DCF_GRID_WACC):
    r = g0 + 2 + i; put(d, f"A{r}", f"={I['wacc']}+{wa}", '0.00%')
    for j in range(len(A.DCF_GRID_G)):
        cc = CL(2 + j); w = f"$A{r}"; g = f"{cc}${g0 + 1}"
        pv = f"SUMPRODUCT(${FC}${DR['ufcf']}:${LC}${DR['ufcf']},1/(1+{w})^${FC}${DR['t']}:${LC}${DR['t']})"
        tv = f"${LC}${DR['ufcf']}*(1+{g})/({w}-{g})/(1+{w})^{N}"
        put(d, f"{cc}{r}", f"=({pv}+{tv}+{I['netcash']})/{I['fd']}*100/{I['fx']}", P0)
DG = {"low": f"DCF!$C${g0 + 2 + 4}", "high": f"DCF!$E${g0 + 2}"}   # WACC+1% & g-0.5% ; WACC-1% & g+0.5%
d.column_dimensions["A"].width = 40

# =============================================================== Comps
cp = S["Comps"]
title(cp, "Listed cybersecurity peers at the undisturbed date",
      "Prices: week ending 19 Apr 2024 (PX). Balance sheets and guidance: latest release before 25 Apr 2024. NTM = next-FY guidance midpoint.")
heads = ["Company", "Ticker", "Price $", "Diluted shares m", "Cash & inv. $m", "Debt $m", "NTM revenue $m", "Last FY revenue $m",
         "Guide year", "Market cap $m", "EV $m", "EV / NTM revenue", "NTM growth", "Source"]
for j, t in enumerate(heads, start=1): hdr(cp, 4, j, t)
for i, (nm, tk, px, sh, cash, debt, ntm, last, gy, sk) in enumerate(A.COMPS, start=5):
    cp[f"A{i}"], cp[f"B{i}"], cp[f"I{i}"], cp[f"N{i}"] = nm, tk, gy, src_label(sk)
    for col, v, fmt in (("C", px, USD), ("D", sh, M1), ("E", cash, M1), ("F", debt, M1), ("G", ntm, M1), ("H", last, M1)):
        put(cp, f"{col}{i}", v, fmt)
    put(cp, f"J{i}", f"=C{i}*D{i}", M); put(cp, f"K{i}", f"=J{i}-E{i}+F{i}", M)
    put(cp, f"L{i}", f"=K{i}/G{i}", X); put(cp, f"M{i}", f"=G{i}/H{i}-1", PCT)
e = 4 + len(A.COMPS); L = f"$L$5:$L${e}"; Mg = f"$M$5:$M${e}"
stats = [("25th percentile", f"=PERCENTILE({L},0.25)"), ("Median", f"=MEDIAN({L})"), ("75th percentile", f"=PERCENTILE({L},0.75)"),
         ("Regression slope (multiple per unit of growth)", f"=SLOPE({L},{Mg})"), ("Regression intercept", f"=INTERCEPT({L},{Mg})"),
         ("R-squared", f"=RSQ({L},{Mg})")]
for i, (lab, f) in enumerate(stats, start=e + 2):
    cp[f"K{i}"] = lab; put(cp, f"L{i}", f, X if i < e + 5 else "0.00")
q1, med, q3, slope, icpt = (f"Comps!$L${e + 2 + k}" for k in range(5))
b0 = e + 10
cp[f"A{b0}"] = "Applied to Darktrace"; cp[f"A{b0}"].font = BOLD
app = [("Darktrace FY25E revenue (DCF case)", "=DCF!C5", M1),
       ("Darktrace NTM revenue at 25 Apr 2024", f"={I['ntm_w']}*{I['rev24']}+(1-{I['ntm_w']})*B{b0 + 1}", M1),
       ("Darktrace NTM growth (FY25E)", "=DCF!C6", PCT), ("Offer EV / NTM revenue", f"={I['ev']}/B{b0 + 2}", X),
       ("Regression-implied multiple", f"={icpt}+{slope}*B{b0 + 3}", X),
       ("Value per share at 25th pct (p)", f"={pps(f'{q1}*B{b0 + 2}')}", P0),
       ("Value per share at median (p)", f"={pps(f'{med}*B{b0 + 2}')}", P0),
       ("Value per share at 75th pct (p)", f"={pps(f'{q3}*B{b0 + 2}')}", P0),
       ("Regression value, multiple -1x (p)", f"={pps(f'(B{b0 + 5}-1)*B{b0 + 2}')}", P0),
       ("Regression value, multiple +1x (p)", f"={pps(f'(B{b0 + 5}+1)*B{b0 + 2}')}", P0)]
for i, (lab, f, fmt) in enumerate(app, start=b0 + 1):
    cp[f"A{i}"] = lab; put(cp, f"B{i}", f, fmt)
CMP = {k: f"Comps!$B${b0 + n}" for k, n in (("q1", 6), ("med", 7), ("q3", 8), ("reg_lo", 9), ("reg_hi", 10),
                                              ("ntm", 2), ("offer_x", 4), ("reg_x", 5))}
cp.column_dimensions["A"].width = 40; cp.column_dimensions["N"].width = 60; cp.column_dimensions["K"].width = 14

# =============================================================== Precedents
pr = S["Precedents"]
title(pr, "Precedent cybersecurity take-privates", "EV / last full fiscal-year revenue before announcement. Premium on the basis stated by each party.")
heads = ["Target", "Buyer", "Announced", "EV $m", "Last FY revenue $m", "EV / revenue", "Premium", "Premium basis", "Deal source", "Revenue source"]
for j, t in enumerate(heads, start=1): hdr(pr, 4, j, t)
for i, (t, b, dt, ev, rv_, p, bs, s1, s2) in enumerate(A.PRECEDENTS, start=5):
    pr[f"A{i}"], pr[f"B{i}"], pr[f"C{i}"], pr[f"H{i}"], pr[f"I{i}"], pr[f"J{i}"] = t, b, dt, bs, src_label(s1), src_label(s2)
    put(pr, f"D{i}", ev, M); put(pr, f"E{i}", rv_, M1); put(pr, f"F{i}", f"=D{i}/E{i}", X); put(pr, f"G{i}", p, PCT)
e2 = 4 + len(A.PRECEDENTS); F_ = f"$F$5:$F${e2}"; G_ = f"$G$5:$G${e2}"
pr[f"A{e2 + 1}"] = "Darktrace (Thoma Bravo)"; pr[f"A{e2 + 1}"].font = BOLD
put(pr, f"D{e2 + 1}", f"={I['ev']}", M); put(pr, f"E{e2 + 1}", f"={I['ltm_rev']}", M1)
put(pr, f"F{e2 + 1}", f"=D{e2 + 1}/E{e2 + 1}", X); put(pr, f"G{e2 + 1}", f"={OFR['Premium to undisturbed close (517p)']}", PCT)
pr[f"H{e2 + 1}"] = "LTM revenue to Dec 2023; premium to close 25 Apr 2024"
st = e2 + 3
for k, (lab, q) in enumerate((("25th percentile", 0.25), ("Median", 0.5), ("75th percentile", 0.75))):
    r = st + k; pr[f"C{r}"] = lab
    put(pr, f"F{r}", f"=PERCENTILE({F_},{q})", X); put(pr, f"G{r}", f"=PERCENTILE({G_},{q})", PCT)
a0 = st + 4; pr[f"A{a0}"] = "Applied to Darktrace"; pr[f"A{a0}"].font = BOLD
LR = I["ltm_rev"]
app = [("EV/revenue on LTM revenue, 25th pct (p)", f"={pps(f'$F${st}*{LR}')}"),
       ("EV/revenue on LTM revenue, median (p)", f"={pps(f'$F${st + 1}*{LR}')}"),
       ("EV/revenue on LTM revenue, 75th pct (p)", f"={pps(f'$F${st + 2}*{LR}')}"),
       ("Premium on undisturbed 517p, 25th pct (p)", f"={I['undisturbed']}*(1+$G${st})"),
       ("Premium on undisturbed 517p, median (p)", f"={I['undisturbed']}*(1+$G${st + 1})"),
       ("Premium on undisturbed 517p, 75th pct (p)", f"={I['undisturbed']}*(1+$G${st + 2})"),
       ("Premium on 3-month VWAP 429.9p, median (p)", f"={I['vwap3m']}*(1+$G${st + 1})")]
for i, (lab, f) in enumerate(app, start=a0 + 1):
    pr[f"A{i}"] = lab; put(pr, f"B{i}", f, P0)
PRC = {k: f"Precedents!$B${a0 + n}" for k, n in (("q1", 1), ("med", 2), ("q3", 3), ("pq1", 4), ("pmed", 5), ("pq3", 6), ("pvwap", 7))}
pr.column_dimensions["A"].width = 42; pr.column_dimensions["H"].width = 30
pr.column_dimensions["I"].width = 50; pr.column_dimensions["J"].width = 50


# =============================================================== LBO blocks
def lbo_block(ws, top, debt_ref, ck, exit_refs, label):
    """Sources & uses, FY25-FY29 operating and debt schedule, and returns for each exit multiple. Mirrors engine.lbo."""
    sh = ws.title
    def q(ref): return f"'{sh}'!{ref}"
    ws[f"A{top}"] = label; ws[f"A{top}"].font = Font(bold=True, size=12)
    su = top + 1
    lines = [("Total debt at close", f"={debt_ref}", M), ("First-lien term loan", f"=B{su}*{I['tlb']}/({I['tlb']}+{I['sl']})", M),
             ("Second-lien term loan", f"=B{su}-B{su + 1}", M), ("Equity purchase price", f"={I['offer_usd']}*{I['fd']}", M),
             ("Entry enterprise value", f"=B{su + 3}-{I['netcash']}", M),
             ("OID", f"=B{su + 1}*(1-{I['tlb_oid']})+B{su + 2}*(1-{I['sl_oid']})", M1),
             ("Financing fees", f"=B{su}*{I['fin_fee']}", M1), ("Advisory and deal fees", f"=B{su + 4}*{I['adv_fee']}", M1),
             ("Total uses", f"=B{su + 3}+B{su + 5}+B{su + 6}+B{su + 7}", M),
             ("Target balance sheet cash used", f"={I['cash_close']}-{I['min_cash']}", M1),
             ("Sponsor equity (plug)", f"=B{su + 8}-B{su}-B{su + 9}", M),
             ("Sponsor equity % of uses", f"=B{su + 10}/B{su + 8}", PCT),
             ("Debt / LTM adjusted EBITDA (Dec 2023)", f"=B{su}/{I['ltm_ebitda']}", X),
             ("Debt / FY24 ARR", f"=B{su}/{I['arr24']}", X)]
    for i, (lab_, f, fmt) in enumerate(lines):
        ws[f"A{su + i}"] = lab_; put(ws, f"B{su + i}", f, fmt, lab_ in ("Total uses", "Sponsor equity (plug)"))
    TLB0, SL0, SPONSOR = f"$B${su + 1}", f"$B${su + 2}", f"$B${su + 10}"
    s0 = su + len(lines) + 1
    keys = ["rev", "g", "ebitda", "m", "capex", "dnwc", "da", "stub", "sofr", "tlb_open", "int_tlb", "int_sl", "interest",
            "deductible", "tax", "fcf", "amort", "cash_open", "avail", "sweep", "tlb", "sl", "cash", "net_debt", "lev", "cover"]
    lab = {"rev": "Revenue", "g": "Growth", "ebitda": "Adjusted EBITDA", "m": "Margin", "capex": "Capex",
           "dnwc": "Working capital inflow", "da": "D&A (tax)", "stub": "Ownership fraction of year", "sofr": "SOFR",
           "int_tlb": "Interest: first lien", "int_sl": "Interest: second lien", "interest": "Total cash interest",
           "deductible": "Deductible interest (30% of tax-EBITDA cap)", "tax": "Cash tax",
           "fcf": "Free cash flow after interest and tax", "tlb_open": "First lien: opening",
           "amort": "First lien: mandatory amortisation", "cash_open": "Opening cash",
           "avail": "Cash available after amortisation", "sweep": "First lien: cash sweep", "tlb": "First lien: closing",
           "sl": "Second lien: closing", "cash": "Closing cash", "net_debt": "Net debt", "lev": "Total debt / EBITDA",
           "cover": "EBITDA / cash interest"}
    R = {k: s0 + 1 + n for n, k in enumerate(keys)}
    hdr(ws, s0, 1, "$m"); hdr(ws, s0, 2, "FY24A")
    for j, y in enumerate(A.YEARS_LBO): hdr(ws, s0, 3 + j, f"FY{y % 100}E")
    for k, r in R.items(): ws[f"A{r}"] = lab[k]
    put(ws, f"B{R['rev']}", "=History!D5", M1)
    for j in range(len(A.YEARS_LBO)):
        c, p = CL(3 + j), CL(2 + j)
        f = {
            "g": f"={I['g_' + ck][j]}", "rev": f"={p}{R['rev']}*(1+{c}{R['g']})", "m": f"={I['m_' + ck][j]}",
            "ebitda": f"={c}{R['rev']}*{c}{R['m']}", "capex": f"={c}{R['rev']}*{I['capex_pct']}",
            "dnwc": f"=({c}{R['rev']}-{p}{R['rev']})*{I['nwc']}", "da": f"={c}{R['rev']}*{I['da_pct']}",
            "stub": f"={I['stub']}" if j == 0 else "=1", "sofr": f"={I['sofr'][j]}",
            "tlb_open": f"={TLB0}" if j == 0 else f"={p}{R['tlb']}",
            "int_tlb": f"={c}{R['tlb_open']}*({c}{R['sofr']}+{I['tlb_m']})*{c}{R['stub']}",
            "int_sl": f"={c}{R['sl']}*({c}{R['sofr']}+{I['sl_m']})*{c}{R['stub']}",
            "interest": f"={c}{R['int_tlb']}+{c}{R['int_sl']}",
            "deductible": f"=MIN({c}{R['interest']},{I['cir']}*{c}{R['ebitda']}*{c}{R['stub']})",
            "tax": f"=MAX(0,({c}{R['ebitda']}-{c}{R['da']})*{c}{R['stub']}-{c}{R['deductible']})*{I['tax']}",
            "fcf": f"=({c}{R['ebitda']}-{c}{R['capex']}+{c}{R['dnwc']}-{I['lease']})*{c}{R['stub']}-{c}{R['interest']}-{c}{R['tax']}",
            "amort": f"=MIN({c}{R['tlb_open']},{TLB0}*{I['tlb_am']}*{c}{R['stub']})",
            "cash_open": f"={I['min_cash']}" if j == 0 else f"={p}{R['cash']}",
            "avail": f"={c}{R['cash_open']}+{c}{R['fcf']}-{c}{R['amort']}",
            "sweep": f"=MIN(MAX(0,{c}{R['avail']}-{I['min_cash']})*{I['sweep']},{c}{R['tlb_open']}-{c}{R['amort']})",
            "tlb": f"={c}{R['tlb_open']}-{c}{R['amort']}-{c}{R['sweep']}", "sl": f"={SL0}",
            "cash": f"={c}{R['avail']}-{c}{R['sweep']}", "net_debt": f"={c}{R['tlb']}+{c}{R['sl']}-{c}{R['cash']}",
            "lev": f"=({c}{R['tlb']}+{c}{R['sl']})/{c}{R['ebitda']}",
            "cover": f"={c}{R['ebitda']}*{c}{R['stub']}/{c}{R['interest']}",
        }
        for k, v in f.items():
            fmt = PCT if k in ("g", "m") else '0.00%' if k == "sofr" else "0.00" if k == "stub" else X if k in ("lev", "cover") else M1
            put(ws, f"{c}{R[k]}", v, fmt, k in ("fcf", "net_debt"))
    e0 = R["cover"] + 2; LCc = CL(2 + len(A.YEARS_LBO))
    hdr(ws, e0, 1, "Returns at exit (30 Jun 2029)")
    for j, h_ in enumerate(["Exit multiple", "Exit EV", "Net debt incl. leases", "Exit equity", "Sponsor share (after MIP)", "MOIC", "IRR"], start=2):
        hdr(ws, e0, j, h_)
    out = []
    for i, xref in enumerate(exit_refs):
        r = e0 + 1 + i
        put(ws, f"B{r}", f"={xref}", X); put(ws, f"C{r}", f"=B{r}*{LCc}{R['ebitda']}", M)
        put(ws, f"D{r}", f"={LCc}{R['net_debt']}+{I['leases_exit']}", M); put(ws, f"E{r}", f"=C{r}-D{r}", M)
        put(ws, f"F{r}", f"=E{r}*(1-{I['mip']})", M); put(ws, f"G{r}", f"=F{r}/{SPONSOR}", X)
        put(ws, f"H{r}", f"=G{r}^(1/{I['hold']})-1", PCT, True)
        out.append({"irr": q(f"$H${r}"), "moic": q(f"$G${r}"), "sponsor_exit": q(f"$F${r}")})
    lo = e0 + 2 + len(exit_refs)
    ws[f"A{lo}"] = "Liquidity low point (minimum closing cash)"; put(ws, f"B{lo}", f"=MIN(C{R['cash']}:{LCc}{R['cash']})", M1)
    return dict(R=R, sponsor=q(SPONSOR), out=out, end=lo, su=su, last=LCc, lev_ltm=q(f"$B${su + 12}"),
                oid=q(f"$B${su + 5}"), fin=q(f"$B${su + 6}"), cash_used=q(f"$B${su + 9}"), low=q(f"$B${lo}"))


lb = S["LBO"]; dn = S["Downside"]; gr = S["LBO grid"]
title(lb, "LBO: Base case", "Thoma Bravo's actual financing; operating case is an assumption. Interest on opening balances (no circularity).")
title(dn, "LBO: Downside case", "Same financing, slower growth and a weaker cost programme")
lb["J3"] = "Actual debt at close ($m)"; put(lb, "K3", f"={I['tlb']}+{I['sl']}", M)
lb["J4"] = "Exit multiples for the table"; lb["J4"].font = BOLD
for k, x in enumerate(A.GRID_EXIT): put(lb, f"K{5 + k}", x, X)
BASE = lbo_block(lb, 4, "LBO!$K$3", "b", [I["x_b"]] + [f"LBO!$K${5 + k}" for k in range(len(A.GRID_EXIT))], "Base case")
DOWN = lbo_block(dn, 4, "LBO!$K$3", "d", [I["x_d"], I["x_b"]], "Downside case")
R_ = BASE["R"]; LCb = BASE["last"]; su_ = BASE["su"]; xr = BASE["end"] - 2 - len(A.GRID_EXIT)
br = BASE["end"] + 2
lb[f"A{br}"] = "Sponsor value bridge, Base case ($m)"; lb[f"A{br}"].font = BOLD
bl = [("Sponsor equity in", f"=$B${su_ + 10}"),
      ("EBITDA growth at entry multiple", f"=({LCb}{R_['ebitda']}-{I['ltm_ebitda']})*($B${su_ + 4}/{I['ltm_ebitda']})"),
      ("Multiple change (34x LTM to exit)", f"=($B${xr}-$B${su_ + 4}/{I['ltm_ebitda']})*{LCb}{R_['ebitda']}"),
      ("Net debt paid down", f"=($B${su_}-{I['min_cash']}+{I['leases_dec23']})-({LCb}{R_['net_debt']}+{I['leases_exit']})"),
      ("Fees and OID", f"=-($B${su_ + 5}+$B${su_ + 6}+$B${su_ + 7})"),
      ("Management incentive plan", f"=-$E${xr}*{I['mip']}"),
      ("Sponsor exit proceeds", f"=SUM(B{br + 1}:B{br + 6})"), ("Check vs returns table (should be 0)", f"=B{br + 7}-$F${xr}")]
for i, (lab_, f) in enumerate(bl, start=br + 1):
    lb[f"A{i}"] = lab_; put(lb, f"B{i}", f, M, lab_ == "Sponsor exit proceeds")
BRIDGE = f"LBO!$B${br + 8}"
for ws in (lb, dn): ws.column_dimensions["A"].width = 44; ws.column_dimensions["J"].width = 26

title(gr, "LBO grid: IRR by exit multiple and debt at close (Base case)", "Each column is a full re-run of the debt schedule (blocks below)")
put(gr, "A3", "Total debt at close ($m)", bold=True); put(gr, "A4", "Debt / LTM EBITDA", bold=True)
for j, dv in enumerate(A.GRID_DEBT):
    put(gr, f"{CL(2 + j)}3", dv, M); put(gr, f"{CL(2 + j)}4", f"={CL(2 + j)}3/{I['ltm_ebitda']}", '0.0"x"')
gr["A5"] = "IRR: exit multiple (rows) x debt (columns)"; gr["A5"].font = BOLD
for i, x in enumerate(A.GRID_EXIT): put(gr, f"A{6 + i}", x, X)
mg = 6 + len(A.GRID_EXIT) + 1
gr[f"A{mg}"] = "Sponsor equity cheque ($m)"; gr[f"A{mg}"].font = BOLD
top = mg + 3; blocks = []
for j, dv in enumerate(A.GRID_DEBT):
    blk = lbo_block(gr, top, f"'LBO grid'!${CL(2 + j)}$3", "b", [f"'LBO grid'!$A${6 + i}" for i in range(len(A.GRID_EXIT))],
                    f"Block {j + 1}: debt ${dv:,.0f}m")
    blocks.append(blk); top = blk["end"] + 3
    for i in range(len(A.GRID_EXIT)):
        put(gr, f"{CL(2 + j)}{6 + i}", "=" + blk["out"][i]["irr"].split("!")[1], PCT)
    put(gr, f"{CL(2 + j)}{mg}", "=" + blk["sponsor"].split("!")[1], M)
gr.column_dimensions["A"].width = 44

# =============================================================== Ability to pay
ap = S["Ability to pay"]
title(ap, "LBO ability to pay", "Exit equity does not depend on the entry price, so the price is solved in closed form for each IRR hurdle")
ap["A4"] = "Sponsor exit proceeds, Base case ($m)"; put(ap, "B4", "=" + BASE["out"][0]["sponsor_exit"], M)
for j, hurdle in enumerate([0.20, 0.25, "pe"]):
    r = 6 + j * 4
    ap[f"A{r}"] = "IRR hurdle"; put(ap, f"B{r}", hurdle if hurdle != "pe" else f"={I['pe_irr']}", PCT)
    ap[f"A{r + 1}"] = "Sponsor equity that clears the hurdle ($m)"; put(ap, f"B{r + 1}", f"=B4/(1+B{r})^{I['hold']}", M)
    ap[f"A{r + 2}"] = "Maximum price per share ($)"
    put(ap, f"B{r + 2}", f"=(B{r + 1}+LBO!$K$3+{BASE['cash_used']}-{BASE['oid']}-{BASE['fin']}+{I['adv_fee']}*{I['netcash']})"
                         f"/({I['fd']}*(1+{I['adv_fee']}))", USD)
    ap[f"A{r + 3}"] = "Maximum price per share (p)"; put(ap, f"B{r + 3}", f"=B{r + 2}*100/{I['fx']}", P0, True)
ATP = {"20": "'Ability to pay'!$B$9", "25": "'Ability to pay'!$B$13", "pe": "'Ability to pay'!$B$17"}
ap["A19"] = "Offer price (p)"; put(ap, "B19", f"={I['offer_p']}", P0)
ap["A20"] = "Base-case IRR at the offer price"; put(ap, "B20", "=" + BASE["out"][0]["irr"], PCT)
ap.column_dimensions["A"].width = 46

# =============================================================== Rival bid
rv = S["Rival bid"]
title(rv, "Would a rival bid make sense?")
lines = [("Strategic: FY25E revenue ($m)", "=DCF!C5", M1), ("Run-rate cost synergies ($m)", f"=B4*{I['syn_pct']}", M1),
         ("Capitalised after-tax synergies less integration cost ($m)", f"=B5*(1-{I['tax']})*{I['syn_mult']}-B5*{I['syn_cost']}", M1),
         ("Standalone DCF EV ($m)", "=DCF!B26", M1), ("EV with share of synergies paid away ($m)", f"=B7+B6*{I['syn_share']}", M1),
         ("Strategic maximum value per share (p)", f"={pps('B8')}", P0),
         ("Rival sponsor at its IRR hurdle, same plan and debt (p)", f"={ATP['pe']}", P0),
         ("Offer (p)", f"={I['offer_p']}", P0),
         ("Strategic headroom vs offer", "=B9/B11-1", PCT), ("Sponsor headroom vs offer", "=B10/B11-1", PCT),
         ("Share of capital not bound by irrevocables", f"=1-{I['irrev']}", PCT)]
for i, (lab_, f, fmt) in enumerate(lines, start=4):
    rv[f"A{i}"] = lab_; put(rv, f"B{i}", f, fmt)
rv["A16"] = ("Irrevocables from directors, KKR DA and Summit Partners stay binding if a higher competing offer is made (R27). "
             "A scheme needs 75% by value of shares voted, so 14.7% committed to Thoma Bravo is a large block against an interloper.")
rv.column_dimensions["A"].width = 62

# =============================================================== Football field
ff = S["Football field"]
title(ff, "Valuation football field (pence per share)", "Offer 620p. Undisturbed close 517p.")
for j, t in enumerate(["Method", "Low", "High", "Width", "Offer"], start=1): hdr(ff, 4, j, t)
items = [("March placing to 12-month high", f"={I['placing']}", f"={I['high12m']}"),
         ("DCF (WACC +/-1%, g +/-0.5%)", f"={DG['low']}", f"={DG['high']}"),
         ("Trading comps EV/NTM revenue (Q1-Q3)", f"={CMP['q1']}", f"={CMP['q3']}"),
         ("Growth-adjusted comps (regression +/-1x)", f"={CMP['reg_lo']}", f"={CMP['reg_hi']}"),
         ("Precedent take-privates EV/LTM revenue (Q1-Q3)", f"={PRC['q1']}", f"={PRC['q3']}"),
         ("Precedent premiums on undisturbed 517p (Q1-Q3)", f"={PRC['pq1']}", f"={PRC['pq3']}"),
         ("LBO ability to pay (25% to 20% IRR)", f"={ATP['25']}", f"={ATP['20']}")]
for i, (lab_, lo, hi) in enumerate(items, start=5):
    ff[f"A{i}"] = lab_; put(ff, f"B{i}", lo, P0); put(ff, f"C{i}", hi, P0); put(ff, f"D{i}", f"=C{i}-B{i}", P0)
    put(ff, f"E{i}", f"={I['offer_p']}", P0)
ch = BarChart(); ch.type = "bar"; ch.grouping = "stacked"; ch.overlap = 100
ch.title = "Darktrace: value per share (p) vs 620p offer"; n_ = 4 + len(items)
ch.add_data(Reference(ff, min_col=2, min_row=4, max_row=n_), titles_from_data=True)
ch.add_data(Reference(ff, min_col=4, min_row=4, max_row=n_), titles_from_data=True)
ch.set_categories(Reference(ff, min_col=1, min_row=5, max_row=n_))
ch.series[0].graphicalProperties.noFill = True; ch.series[0].graphicalProperties.line.noFill = True
ch.series[1].graphicalProperties.solidFill = "1F3864"; ch.height, ch.width = 9, 18; ch.legend = None
ff.add_chart(ch, "G4"); ff.column_dimensions["A"].width = 50

# =============================================================== Sources and Cover
src = S["Sources"]; title(src, "Sources")
hdr(src, 3, 1, "Key"); hdr(src, 3, 2, "Document"); hdr(src, 3, 3, "URL")
for i, (k, (t, u)) in enumerate(A.SOURCES.items(), start=4): src[f"A{i}"], src[f"B{i}"], src[f"C{i}"] = k, t, u
src[f"A{5 + len(A.SOURCES)}"] = "Conventions: blue = input, black = formula, green = link to another sheet. Not investment advice."
src.column_dimensions["B"].width = 80; src.column_dimensions["C"].width = 90

title(cover, "Darktrace plc / Thoma Bravo take-private (2024)")
cover["A2"] = ("Project 04: sell-side fairness view. Project 08: LBO and IC case. "
               "Retrospective analysis: April market inputs, July FY24 update; not an April information set.")
cover["A3"] = "Tobenna Agubuche | github.com/Tobenna30/finance-models | Not investment advice"
cov = [("Offer price (p)", f"={I['offer_p']}", P0), ("Premium to undisturbed close", f"={OFR['Premium to undisturbed close (517p)']}", PCT),
       ("Premium to 3-month VWAP", f"={OFR['Premium to 3-month VWAP (429.9p)']}", PCT),
       ("Offer EV / LTM revenue", f"={OFR['EV / LTM revenue']}", X), ("Offer EV / NTM revenue", f"={CMP['offer_x']}", X),
       ("DCF value per share (p)", f"={DCFP}", P0), ("Peer median EV/NTM value (p)", f"={CMP['med']}", P0),
       ("Precedent median EV/LTM value (p)", f"={PRC['med']}", P0), ("Precedent median premium value (p)", f"={PRC['pmed']}", P0),
       ("LBO ability to pay at 20% IRR (p)", f"={ATP['20']}", P0), ("Strategic maximum (p)", "='Rival bid'!B9", P0),
       ("LBO base IRR", "=" + BASE["out"][0]["irr"], PCT), ("LBO base MOIC", "=" + BASE["out"][0]["moic"], X),
       ("LBO downside IRR", "=" + DOWN["out"][0]["irr"], PCT), ("LBO downside MOIC", "=" + DOWN["out"][0]["moic"], X),
       ("Sponsor equity cheque ($m)", "=" + BASE["sponsor"], M), ("Debt / LTM EBITDA at close", "=" + BASE["lev_ltm"], X),
       ("Base liquidity low point ($m)", "=" + BASE["low"], M1),
       ("Base shortfall to operating cash minimum ($m)", f"=MAX(0,{I['min_cash']}-{BASE['low']})", M1),
       ("Downside liquidity low point ($m)", "=" + DOWN["low"], M1),
       ("Downside shortfall to operating cash minimum ($m)", f"=MAX(0,{I['min_cash']}-{DOWN['low']})", M1), ("Bridge check (should be 0)", "=" + BRIDGE, M1)]
for i, (lab_, f, fmt) in enumerate(cov, start=5):
    cover[f"A{i}"] = lab_; put(cover, f"B{i}", f, fmt)
cover.column_dimensions["A"].width = 58; cover.column_dimensions["B"].width = 14
for ws in wb.worksheets: ws.sheet_view.showGridLines = False

if __name__ == "__main__":
    RAW.parent.mkdir(parents=True, exist_ok=True)
    wb.save(RAW)
    print("saved", RAW)
