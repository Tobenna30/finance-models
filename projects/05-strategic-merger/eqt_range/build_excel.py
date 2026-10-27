"""Openpyxl authoring follows the repository's explicit house standard."""
import csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import ColorScaleRule, CellIsRule
from openpyxl.chart import BarChart, Reference
from openpyxl.workbook.properties import CalcProperties
from .assumptions import ROOT, INPUTS, SOURCES, SENS_RATIOS, SENS_SYNERGIES

RAW = ROOT/'models'/'EQT_Range_Merger_raw.xlsx'
OUT = ROOT/'models'/'EQT_Range_Merger.xlsx'
I = {key: i for i,key in enumerate(INPUTS, 6)}
ACTIVE = len(INPUTS)+8
NUM = '#,##0.0;(#,##0.0);"-"'
PCT = '0.0%;(0.0%);"-"'
EPS = '0.00;(0.00);"-"'
NAVY = '1F3864'

def inp(key):
    return f'Inputs!$D${I[key]}'

def main():
    wb = Workbook(); wb.remove(wb.active)
    for name in ('Cover','Inputs','Merger','Sensitivity','Value','History','Sources'):
        s = wb.create_sheet(name)
        s.sheet_view.showGridLines = False
        for col,width in {'A':2,'B':2,'C':53,'D':17,'E':17,'F':17,'G':65,'H':17,'I':17}.items():
            s.column_dimensions[col].width = width
        s.sheet_properties.pageSetUpPr.fitToPage = True
        s.page_setup.orientation = 'landscape'; s.page_setup.paperSize = s.PAPERSIZE_A3
        s.page_setup.fitToWidth=1; s.page_setup.fitToHeight=0
        s.freeze_panes='D6' if name not in ('Cover','Value') else None
        s['C2']=name if name!='Cover' else 'EQT / Range Resources'
        s['C2'].font=Font(name='Arial',size=16,bold=True,color=NAVY)
        s['C3']='Announcement: 26 Oct 2026. Earnings: LTM Jun 2026. USDm / million shares.'
        s['C3'].font=Font(name='Arial',size=10,italic=True)
        s.merge_cells('C3:I3')
    def put(s, row, label, vals, fmt=NUM):
        s.cell(row,3,label)
        for col,v in enumerate(vals,4):
            c=s.cell(row,col,v); c.number_format=fmt
    def note(s,row,text):
        s.cell(row,7,text); s.cell(row,7).alignment=Alignment(wrap_text=True,vertical='top')
        s.row_dimensions[row].height=max(22,14*(len(text)//65+1))
    def header(s,row,labels):
        for col,v in enumerate(labels,3):
            c=s.cell(row,col,v); c.fill=PatternFill('solid',fgColor=NAVY)
            c.font=Font(name='Arial',size=10,color='FFFFFF',bold=True)
    s=wb['Inputs']; header(s,5,['Input','Value','Units','Tag','Source / reasoning'])
    for key,(label,value,tag,unit,why) in INPUTS.items():
        r=I[key];put(s,r,label,[value]);s.cell(r,5,unit);s.cell(r,6,tag);note(s,r,why)
        s.cell(r,4).fill=PatternFill('solid',fgColor='FFF2CC')
        if unit=='%': s.cell(r,4).number_format=PCT
        elif unit in ('USD/share','EQT/RRC'):s.cell(r,4).number_format='0.000'
    put(s,ACTIVE,'Active Range earnings multiplier',[f'=CHOOSE({inp("case")},{inp("r_down")},{inp("r_base")},{inp("r_up")})'])
    put(s,ACTIVE+1,'Active synergy target realised',[f'=CHOOSE({inp("case")},{inp("s_down")},{inp("s_base")},{inp("s_up")})'],PCT)
    put(s,ACTIVE+2,'Case selected',[f'=CHOOSE({inp("case")},"Down","Base","Up")'])
    dv=DataValidation(type='whole',operator='between',formula1=1,formula2=3);dv.error='Choose 1, 2 or 3';dv.showErrorMessage=True;dv.errorTitle='Invalid case';dv.errorStyle='stop';s.add_data_validation(dv);dv.add(s.cell(I['case'],4))
    for name in ('Cover','Merger','Sensitivity','Value','History'):
        wb[name]['H2']='Case selected';wb[name]['I2']=f'=Inputs!D{ACTIVE+2}'
    # Rebuild LTM from matching accounting measures; do not sum per-share values.
    s=wb['History'];header(s,5,['Earnings measure','FY2025','H1 2025','H1 2026','LTM Jun 2026'])
    for r,key in enumerate(('e_adj','e_gaap','r_adj','r_gaap','r_sbc'),6):
        label={'e_adj':'EQT adjusted attributable NI','e_gaap':'EQT GAAP attributable NI','r_adj':'Range adjusted NI','r_gaap':'Range GAAP NI','r_sbc':'Range SBC addback, before tax'}[key]
        put(s,r,label,[f'={inp(key+suffix)}' for suffix in ('25','125','126')]+[f'=D{r}-E{r}+F{r}'])
    note(s,12,'LTM = FY2025 minus H1 2025 plus H1 2026. Q2 diluted shares are closing proxies, not historical LTM EPS denominators.')
    note(s,14,'Range SBC uses the net income reconciliation, including amortisation. EQT adjusted reconciliation has no SBC addback. Other adjustment policies still differ.')
    # Shared economics above the three full post-close years.
    s=wb['Merger'];header(s,5,['Transaction / earnings bridge','Value'])
    rows={
      6:('EQT standalone normalised earnings','=History!G6'),
      7:('Range adjusted earnings','=History!G8'),
      8:('Range SBC retained, after tax',f'=History!G10*(1-{inp("tax")})'),
      9:('Range normalised earnings, active case',f'=(D7-D8)*Inputs!D{ACTIVE}'),
      10:('EQT diluted shares, closing proxy',f'={inp("eqt_shares")}'),
      11:('New EQT shares, exchange proxy',f'={inp("rrc_shares")}*{inp("ratio")}'),
      12:('Combined diluted shares','=SUM(D10:D11)'),
      13:('Range ownership','=D11/D12'),
      14:('EQT standalone normalised EPS','=D6/D10'),
      15:('Annual incremental DD&A',f'={inp("stepup")}/{inp("life")}'),
      16:('No-synergy recurring accretion',f'=((D6+D9-D15*(1-{inp("tax")}))/D12)/D14-1'),
      17:('EPS break-even annual pre-tax synergies',f'=(D14*D11-D9)/(1-{inp("tax")})+D15'),
    }
    for r,(label,f) in rows.items():put(s,r,label,[f],PCT if r in (13,16) else EPS if r==14 else NUM)
    note(s,6,'Normalised lens retains Range stock compensation at an assumed 25% tax rate. Other company non-GAAP adjustments are retained.')
    note(s,11,'Actual closing shares and award conversion remain unverified. Shares issued are not the disclosed final number.')
    note(s,15,'Illustrative purchase accounting: USD1,500m step-up / 15 years. Switch step-up to zero to isolate the impact.')
    header(s,20,['Full post-close years','Year 1','Year 2','Year 3'])
    labels={21:'Average synergy capture',22:'Pre-tax synergies',23:'Integration costs, cash and P&L',24:'Transaction fees, P&L',25:'Normalised recurring net income',26:'Normalised recurring EPS',27:'Normalised recurring accretion',28:'Net income including one-off costs',29:'EPS including one-off costs',30:'Accretion including one-off costs',31:'Company-adjusted EPS, before PPA',32:'Company-adjusted accretion, before PPA',33:'GAAP earnings lens incl costs, EPS',34:'GAAP earnings lens incl costs, accretion'}
    for r,label in labels.items():s.cell(r,3,label)
    for col in 'DEF':
        ramp=f'={inp("ramp1")}' if col=='D' else '=1'
        cost=f'={inp("integration")}*{inp("cost1")}' if col=='D' else f'={inp("integration")}*(1-{inp("cost1")})' if col=='E' else '=0'
        formulas={21:ramp,22:f'={inp("synergy")}*Inputs!$D${ACTIVE+1}*{col}21',23:cost,24:f'={inp("fees")}' if col=='D' else '=0',25:f'=$D$6+$D$9+({col}22-$D$15)*(1-{inp("tax")})',26:f'={col}25/$D$12',27:f'={col}26/$D$14-1',28:f'={col}25-{col}23*(1-{inp("tax")})-{col}24',29:f'={col}28/$D$12',30:f'={col}29/$D$14-1',31:f'=($D$6+$D$7*Inputs!$D${ACTIVE}+{col}22*(1-{inp("tax")}))/$D$12',32:f'={col}31/$D$14-1',33:f'=(History!$G$7+History!$G$9*Inputs!$D${ACTIVE}+({col}22-$D$15-{col}23)*(1-{inp("tax")} )-{col}24)/$D$12',34:f'={col}33/(History!$G$7/$D$10)-1'}
        for r,f in formulas.items():s[f'{col}{r}']=f;s[f'{col}{r}'].number_format=PCT if r in (21,27,30,32,34) else EPS if r in (26,29,31,33) else NUM
    note(s,28,'One-offs reduce earnings and cash. No acquisition cash payment or new financing interest: stock consideration only; existing debt interest remains in earnings.')
    note(s,33,'Uses historical GAAP NI plus illustrative adjustments. Not GAAP pro forma EPS: final PPA, deferred taxes, debt fair values and closing stub unavailable.')
    # Value test pays the premium in stock, which still has an economic cost.
    s=wb['Value'];header(s,5,['Value versus equity premium','Value'])
    formulas={
      6:('Implied offer per Range share',f'={inp("ratio")}*{inp("eqt_price")}'),
      7:('Equity consideration, share proxy',f'=D6*{inp("rrc_shares")}'),
      8:('Premium to unaffected reference',f'=D6/{inp("rrc_price")}-1'),
      9:('Equity premium paid',f'=(D6-{inp("rrc_price")})*{inp("rrc_shares")}'),
      10:('After-tax synergy PV per USD1 run rate',f'=(1-{inp("tax")})*({inp("ramp1")}/(1+{inp("discount")})+1/({inp("discount")}*(1+{inp("discount")})))'),
      11:('Gross synergy PV, perpetuity',f'={inp("synergy")}*Inputs!D{ACTIVE+1}*D10'),
      12:('PV of integration and fees',f'={inp("integration")}*(1-{inp("tax")})*({inp("cost1")}/(1+{inp("discount")})+(1-{inp("cost1")})/(1+{inp("discount")})^2)+{inp("fees")}'),
      13:('Net synergy PV','=D11-D12'),
      14:('Value surplus after premium','=D13-D9'),
      15:('Value break-even pre-tax annual synergy','=(D9+D12)/D10'),
      16:('Maximum exchange ratio at value break-even',f'=({inp("rrc_price")}+D13/{inp("rrc_shares")})/{inp("eqt_price")}'),
      17:('Finite 15-year synergy PV',f'={inp("synergy")}*Inputs!D{ACTIVE+1}*(1-{inp("tax")})*((1-(1+{inp("discount")})^-15)/{inp("discount")}-(1-{inp("ramp1")})/(1+{inp("discount")}))'),
      18:('Finite 15-year value surplus','=D17-D12-D9'),
      20:('Stock consideration funding',f'=Merger!D11*{inp("eqt_price")}'),
      21:('Stock consideration funding less purchase price','=D20-D7'),
      22:('Cash used for integration and fees',f'={inp("integration")}+{inp("fees")}'),
      23:('Cash bridge allocation check',f'=SUM(Merger!D23:F23)+Merger!D24-D22'),
    }
    for r,(label,f) in formulas.items():put(s,r,label,[f],PCT if r==8 else '0.000' if r==16 else EPS if r in (6,21,23) else NUM)
    note(s,6,'Reference prices are secondary cross-checks. EQT is a rounded implied proxy. This is an announcement-date price test, not a live quote.')
    note(s,11,'Optimistic perpetuity: no growth, no incremental capex or working capital; 50% year 1, 100% from year 2. No speculative >600m pathway.')
    note(s,14,'Premium paid in shares is an economic cost. EPS accretion alone does not establish value creation. Unaffected share price is a standalone value proxy.')
    note(s,22,'Gross cash cost reserve, before tax benefit. No opening liquidity or refinancing capacity assertion. This module is not a full three-statement model.')
    # A simple two-variable formula is appropriate; no duplicated complex engine.
    s=wb['Sensitivity'];header(s,5,['Year 3 recurring accretion','0','162.5','250','325','450','600'])
    s['C4']='Assumed ratios down / annual pre-tax synergies across (USDm)'
    s.merge_cells('C4:I4')
    for col,syn in enumerate(SENS_SYNERGIES,4):s.cell(5,col,syn)
    for row,ratio in enumerate(SENS_RATIOS,6):
        s.cell(row,3,ratio);s.cell(row,3).number_format='0.000'
        for col in range(4,10):
            letter=s.cell(5,col).column_letter
            s.cell(row,col,f'=(Merger!$D$6+Merger!$D$9+({letter}$5-Merger!$D$15)*(1-{inp("tax")}))/(Merger!$D$10+{inp("rrc_shares")}*$C{row})/Merger!$D$14-1').number_format=PCT
    s.conditional_formatting.add('D6:I10',ColorScaleRule(start_type='min',start_color='F3D5D3',mid_type='num',mid_value=0,mid_color='FFFFFF',end_type='max',end_color='D5E6D8'))
    note(s,13,'Live simple formula: current case earnings and DD&A, absolute synergy amounts. 600m is sensitivity only. Grid omits one-off costs.')
    # Cover: one recommendation and linked results.
    s=wb['Cover'];header(s,5,['Announcement analysis','Result'])
    put(s,6,'Call',[ '=IF(Value!D14<0,"REJECT AS PRICED","ACCEPT ON MODEL ASSUMPTIONS")'])
    for r,label,f,fmt in [(8,'Year 1 EPS effect including costs','=Merger!D30',PCT),(9,'Year 3 recurring EPS effect','=Merger!F27',PCT),(10,'Company-adjusted year 3, before PPA','=Merger!F32',PCT),(11,'EPS break-even annual synergies','=Merger!D17',NUM),(12,'Value break-even annual synergies','=Value!D15',NUM),(13,'Value surplus, perpetuity (USDm)','=Value!D14',NUM),(14,'Value surplus, 15 years (USDm)','=Value!D18',NUM),(15,'Value break-even exchange ratio','=Value!D16','0.000')]:put(s,r,label,[f],fmt)
    s.merge_cells('D6:F6')
    note(s,18,'LTM earnings held constant over three full post-close years. These are earnings lenses, not calendar forecasts or consensus estimates.')
    note(s,20,'Closing share counts, primary price verification, final purchase accounting and integration plan are review items. Completion subject to approvals, expected Q1 2027.')
    note(s,22,'Change the case on Inputs. Down stresses Range earnings by 20% and captures half the synergy target. EQT earnings stay constant; no implied Henry Hub forecast.')
    ch=BarChart();ch.title='Recurring EPS impact';ch.y_axis.title='Accretion / dilution';ch.add_data(Reference(wb['Merger'],min_col=4,max_col=6,min_row=27,max_row=27),from_rows=True,titles_from_data=False);ch.set_categories(Reference(wb['Merger'],min_col=4,max_col=6,min_row=20,max_row=20));ch.height=8.5;ch.width=14;ch.legend=None;ch.y_axis.scaling.min=-.04;ch.y_axis.scaling.max=.02;ch.y_axis.majorUnit=.02;ch.x_axis.tickLblPos='low';s.add_chart(ch,'G6')
    s=wb['Sources'];header(s,5,['ID','Source and locator','URL'])
    s.column_dimensions['D'].width=80;s.column_dimensions['E'].width=100
    for r,(key,(label,url)) in enumerate(SOURCES.items(),6):
        s.cell(r,3,key);s.cell(r,4,label);s.cell(r,5,url);s.cell(r,5).hyperlink=url
        s.cell(r,4).alignment=Alignment(wrap_text=True);s.cell(r,5).alignment=Alignment(wrap_text=True);s.row_dimensions[r].height=65
    for sheet in wb:
        sheet.sheet_properties.tabColor = NAVY if sheet.title=='Cover' else '6E8BB5' if sheet.title=='Inputs' else None
        sheet.print_options.horizontalCentered=True
        sheet.print_area=f'C2:I{sheet.max_row}'
        for row in sheet:
            for c in row:
                if c.value is None or c.row in (2,3,5,20) and c.fill.fgColor.rgb==NAVY:continue
                if c.font.name is None or c.font.name=='Calibri':
                    formula=c.data_type=='f'
                    colour='000000' if sheet.title=='Cover' else '008000' if formula and '!' in c.value else '000000' if formula else '0000FF' if isinstance(c.value,(int,float)) else '222222'
                    c.font=Font(name='Arial',size=10,color=colour)
                if c.alignment.wrap_text is not True:c.alignment=Alignment(vertical='center',horizontal='right' if isinstance(c.value,(int,float)) or c.data_type=='f' else 'left')
        for r in (6,9,11,12,13,14,15) if sheet.title=='Cover' else ():
            sheet.row_dimensions[r].height=25
        if sheet.title not in ('Sources','History'):
            sheet.conditional_formatting.add('D6:F40',CellIsRule(operator='lessThan',formula=['0'],font=Font(color='9C0006')))
    # Keep source values and their provenance in an easy-to-review machine-readable file.
    (ROOT/'data').mkdir(exist_ok=True)
    with (ROOT/'data'/'inputs.csv').open('w',newline='') as f:
        writer=csv.writer(f,lineterminator='\n');writer.writerow(['key','label','value','tag','unit','source_or_reason'])
        for key,row in INPUTS.items():writer.writerow([key,*row])
    wb.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True)
    RAW.parent.mkdir(exist_ok=True);wb.save(RAW);print(RAW)

if __name__=='__main__':main()
