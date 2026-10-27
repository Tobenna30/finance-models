from pathlib import Path
import math
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

def workbook(sheets):
    w=Workbook();w.remove(w.active)
    for name,rows in sheets.items():
        s=w.create_sheet(name)
        for row in rows:s.append(row)
        s.freeze_panes='B4'
        s.sheet_view.showGridLines=False
        for row in s:
            for c in row:
                c.font=Font(name='Aptos',size=10,color='000000' if c.data_type=='f' else '0000FF' if isinstance(c.value,(int,float)) else '222222')
                c.alignment=Alignment(vertical='top',wrap_text=True)
                if c.data_type=='f' and '!' in c.value:c.font=Font(name='Aptos',size=10,color='008000')
        for c in s[1]:c.fill=PatternFill('solid',fgColor='17365D');c.font=Font(name='Aptos',size=14,color='FFFFFF',bold=True)
        s.row_dimensions[1].height=30
        for i in range(1,s.max_row+1):
            if i>1:s.row_dimensions[i].height=34 if name in ['Sources','Requests'] else 28
        for j in range(1,s.max_column+1):s.column_dimensions[get_column_letter(j)].width=28 if j>1 else 43
        s.sheet_properties.pageSetUpPr.fitToPage=True
        s.page_setup.orientation='landscape';s.page_setup.paperSize=s.PAPERSIZE_A4
        s.page_setup.fitToWidth=1;s.page_setup.fitToHeight=0
        s.print_options.horizontalCentered=True
        if name in ['Cover','Sources']:
            s.column_dimensions['B'].width=105
        for row in s:
            occupied=[c for c in row if c.value is not None]
            if len(occupied)==2 and occupied[0].column==1 and occupied[1].column==2 and isinstance(occupied[1].value,str) and len(occupied[1].value)>60 and s.max_column>2:
                s.merge_cells(start_row=row[0].row,start_column=2,end_row=row[0].row,end_column=s.max_column)
                width=sum(s.column_dimensions[get_column_letter(j)].width for j in range(2,s.max_column+1))
                s.row_dimensions[row[0].row].height=max(28,14*math.ceil(len(occupied[1].value)/(width*.9))+10)
            else:
                h=28
                for c in occupied:
                    if isinstance(c.value,str) and c.data_type!='f':
                        width=s.column_dimensions[c.column_letter].width
                        h=max(h,14*math.ceil(len(c.value)/(width*.85))+10)
                s.row_dimensions[row[0].row].height=h
    return w
