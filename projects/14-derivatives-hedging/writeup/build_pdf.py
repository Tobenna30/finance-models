"""Render a pitch written in simple Markdown to a styled DOCX and PDF.
Supports: # / ## headings, paragraphs, **bold** runs, *italic* paragraphs, - bullets, pipe tables.
Run from this folder: python3 build_pdf.py <file>.md"""
import os, re, subprocess, sys
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY = RGBColor(0x1F, 0x38, 0x64)
md = Path(sys.argv[1]); out = md.with_suffix(".docx")
d = Document()
sec = d.sections[0]
sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
sec.top_margin = sec.bottom_margin = Cm(1.5); sec.left_margin = sec.right_margin = Cm(1.8)
st = d.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(9.5)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
st.paragraph_format.space_after = Pt(4)

def runs(p, text, size=None, italic=False):
    for part in re.split(r"(\*\*[^*]+\*\*)", text):
        if not part: continue
        r = p.add_run(part[2:-2] if part.startswith("**") else part)
        r.bold = part.startswith("**"); r.italic = italic
        if size: r.font.size = Pt(size)

def shade(cell, hex_):
    tcPr = cell._tc.get_or_add_tcPr(); s = OxmlElement("w:shd")
    s.set(qn("w:val"), "clear"); s.set(qn("w:fill"), hex_); tcPr.append(s)

def table(rows):
    rows = [r for r in rows if not re.match(r"^\|[\s:|-]+\|$", r)]
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    n = len(cells[0]); first = 7.5 if n <= 3 else 3.6
    widths = [Cm(first)] + [Cm((17.4 - first) / (n - 1))] * (n - 1)
    t = d.add_table(rows=len(cells), cols=n)
    t.autofit = False
    for i, row in enumerate(cells):
        for j, val in enumerate(row):
            c = t.cell(i, j); c.width = widths[j]
            p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
            if j: p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            runs(p, val, size=8.5)
            if i == 0:
                shade(c, "1F3864")
                for r in p.runs: r.bold = True; r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            elif i % 2 == 0:
                shade(c, "EEF2F8")
    d.add_paragraph().paragraph_format.space_after = Pt(0)

lines = md.read_text().splitlines(); i = 0
while i < len(lines):
    ln = lines[i]
    if ln.startswith("|"):
        block = []
        while i < len(lines) and lines[i].startswith("|"): block.append(lines[i]); i += 1
        table(block); continue
    if ln.startswith("!!"):
        paths = [md.parent / x.strip() for x in ln[2:].split("|")]
        t = d.add_table(rows=1, cols=len(paths)); t.autofit = False
        for j, pth in enumerate(paths):
            cpar = t.cell(0, j).paragraphs[0]; cpar.paragraph_format.space_after = Pt(0)
            cpar.add_run().add_picture(str(pth), width=Cm(17.0 / len(paths)))
        i += 1; continue
    if ln.startswith("# "):
        p = d.add_paragraph(); r = p.add_run(ln[2:]); r.bold = True; r.font.size = Pt(16); r.font.color.rgb = NAVY
        p.paragraph_format.space_after = Pt(2)
    elif ln.startswith("## "):
        p = d.add_paragraph(); r = p.add_run(ln[3:]); r.bold = True; r.font.size = Pt(11); r.font.color.rgb = NAVY
        p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(2)
    elif ln.startswith("- "):
        p = d.add_paragraph(style="List Bullet"); runs(p, ln[2:]); p.paragraph_format.space_after = Pt(2)
    elif ln.startswith("*") and ln.endswith("*") and not ln.startswith("**"):
        p = d.add_paragraph(); runs(p, ln[1:-1], size=8, italic=True)
    elif ln.strip():
        p = d.add_paragraph(); runs(p, ln)
    i += 1
d.save(out)
subprocess.run([os.environ.get("SOFFICE", "soffice"), "--headless", "--convert-to", "pdf", "--outdir", str(md.parent), str(out)], check=True, capture_output=True)
print("built", md.with_suffix(".pdf"))
