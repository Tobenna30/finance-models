"""Repository-standard openpyxl workbook. Forecasts are live formulas."""
import csv
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule,ColorScaleRule
from .assumptions import ROOT,INPUTS,SOURCES,MONTHS,UPSIZE_MONTHS,EV_GRID,EQUITY_GRID

RAW=ROOT/'models/Thames_Restructuring_raw.xlsx';OUT=ROOT/'models/Thames_Restructuring.xlsx'
I={k:r for r,k in enumerate(INPUTS,6)}
NUM='#,##0.0;(#,##0.0);"-"';PCT='0.00%;(0.00%);"-"'
def inp(k):return f'Inputs!$D${I[k]}'

def main():
    w=Workbook();w.remove(w.active)
    for name in ('Cover','Inputs','Liquidity','Recovery','Class B','Sensitivity','Sources'):
        s=w.create_sheet(name);s.sheet_view.showGridLines=False;s.column_dimensions['A'].width=2;s.column_dimensions['B'].width=2;s.column_dimensions['C'].width=55
        for c in range(4,16):s.column_dimensions[__import__('openpyxl').utils.get_column_letter(c)].width=16
        s['C2']='Thames Water restructuring' if name=='Cover' else name;s['C2'].font=Font(name='Arial',size=16,bold=True,color='1F3864')
        s.merge_cells('C3:K3');s['C3']='27 Oct 2026. GBPm. Funding and plan terms are conditional; current cash and claims unverified.';s['C3'].font=Font(name='Arial',size=10,italic=True)
        s.freeze_panes='D6' if name!='Cover' else None;s.sheet_properties.pageSetUpPr.fitToPage=True;s.page_setup.orientation='landscape';s.page_setup.fitToWidth=1;s.page_setup.fitToHeight=0
    def head(s,r,labels):
        for c,v in enumerate(labels,3):s.cell(r,c,v);s.cell(r,c).fill=PatternFill('solid',fgColor='1F3864');s.cell(r,c).font=Font(name='Arial',size=10,bold=True,color='FFFFFF')
    def put(s,r,label,values,fmt=NUM):
        s.cell(r,3,label)
        for c,v in enumerate(values,4):s.cell(r,c,v);s.cell(r,c).number_format=fmt
    def note(s,r,text):
        s.merge_cells(start_row=r,start_column=3,end_row=r,end_column=11);s.cell(r,3,text);s.cell(r,3).alignment=Alignment(wrap_text=True,vertical='center');s.row_dimensions[r].height=max(30,14*(len(text)//140+1))
    s=w['Inputs'];s.column_dimensions['G'].width=85;head(s,5,['Input','Value','Units','Tag','Source / reason'])
    for k,(label,value,tag,unit,why) in INPUTS.items():
        r=I[k];put(s,r,label,[value],PCT if unit in ('%','price/par') else NUM);s.cell(r,5,unit);s.cell(r,6,tag);s.cell(r,7,why);s.cell(r,7).alignment=Alignment(wrap_text=True);s.row_dimensions[r].height=30;s.cell(r,4).fill=PatternFill('solid',fgColor='FFF2CC')
    for key,max_ in [('funding_case',2),('maturity_case',1)]:
        d=DataValidation(type='whole',operator='between',formula1=0,formula2=max_);d.showErrorMessage=True;d.errorStyle='stop';s.add_data_validation(d);d.add(s.cell(I[key],4))
    case_row=len(INPUTS)+8
    put(s,case_row,'Funding case selected',[f'=CHOOSE({inp("funding_case")}+1,"Blocked","Existing only","Upsize assumed")'])
    for name in ('Cover','Liquidity','Recovery','Class B','Sensitivity'):
        w[name]['I2']='Funding case';w[name]['J2']=f'=Inputs!D{case_row}';w[name]['J2'].font=Font(name='Arial',size=10,color='008000')
    s=w['Liquidity'];head(s,5,['Monthly cash use / funding']+[datetime.combine(m,datetime.min.time()) for m in MONTHS])
    for c in range(4,16):s.cell(5,c).number_format='mmm-yy'
    labels={7:'Opening unfunded cash position',8:'Existing capacity draw, conditional',9:'New upsize draw, conditional',10:'Upsize draw fee',11:'Interest on incremental draws',12:'Class A maturity cash requirement',13:'Closing unfunded cash position',14:'Gap to minimum cash buffer',15:'Incremental debt opening',16:'Super senior closing claim',18:'Cumulative incremental draws',19:'Cash bridge check'}
    for r,label in labels.items():s.cell(r,3,label)
    for i in range(12):
        col=s.cell(5,i+4).column_letter;prev=s.cell(5,i+3).column_letter
        fs={7:f'={inp("cash")}' if i==0 else f'={prev}13',8:f'=IF({inp("funding_case")}>=1,{inp("facility")}-{inp("drawn")},0)' if i==0 else '=0',9:f'=IF({inp("funding_case")}=2,{inp("upsize")}/6,0)' if i in UPSIZE_MONTHS else '=0',10:f'={col}9*{inp("fee")}',15:'=0' if i==0 else f'={prev}18',11:f'=({col}15+SUM({col}8:{col}9)/2)*{inp("ss_rate")}/12',12:f'=IF({inp("maturity_case")}=0,{inp("march_claim")},0)' if i==4 else '=0',13:f'={col}7+SUM({col}8:{col}9)-{inp("burn")}-SUM({col}10:{col}12)',14:f'=MAX(0,{inp("buffer")}-{col}13)',18:f'={col}15+SUM({col}8:{col}9)',16:f'={inp("drawn")}+{col}18',19:f'={col}13-({col}7+SUM({col}8:{col}9)-{inp("burn")}-SUM({col}10:{col}12))'}
        for r,f in fs.items():s[f'{col}{r}']=f;s[f'{col}{r}'].number_format='0.00' if r==19 else NUM
    put(s,22,'Minimum unfunded cash position',['=MIN(D13:O13)']);put(s,23,'Peak gap to minimum buffer',['=MAX(D14:O14)']);put(s,24,'Incremental interest paid',['=SUM(D11:O11)']);put(s,25,'Upsize fees paid',['=SUM(D10:O10)'])
    note(s,28,'Negative cash positions show unfinanced need, not actual bank overdrafts. There is no automatic borrowing plug. Existing debt interest is included in monthly cash use. Only incremental draw interest is added, using half-month draw convention.')
    note(s,30,'Six GBP400m upsize tranches are assumptions. The GBP350m remaining existing capacity is counted once, not separately at issuer and TWUL. All draws require approvals and conditions. March maturity stress is separate from funding consent.')
    note(s,32,'No current opening cash verified. Monthly use is an analyst stress. FY26 historical monthly outflow is not a current burn forecast. No full balance sheet is constructed from incomplete debt and asset information.')
    s=w['Recovery'];head(s,5,['Illustrative priority waterfall','Claim / value','Paid','Recovery'])
    put(s,6,'Enterprise value',[f'={inp("ev")}']);put(s,7,'Costs / value leakage',[f'=D6*{inp("cost")}']);put(s,8,'Distributable value',['=D6-D7'])
    put(s,10,'Priority hedge / other reserve',[f'={inp("hedges")}','=MIN(D10,$D$8)','=E10/D10'])
    put(s,11,'Super senior, including assumed new funding',['=Liquidity!O16','=MIN(D11,MAX(0,$D$8-E10))','=E11/D11'])
    put(s,12,'Class A claim proxy',[f'={inp("class_a")}','=MIN(D12,MAX(0,$D$8-SUM(E10:E11)))','=E12/D12'])
    put(s,13,'Class B',[f'={inp("b_claim")}','=MIN(D13,MAX(0,$D$8-SUM(E10:E12)))','=E13/D13'])
    put(s,14,'Residual equity',['=MAX(0,D8-SUM(E10:E13))'])
    put(s,16,'Value allocation check',['=D8-SUM(E10:E13)-D14'],'0.00')
    put(s,18,'Class A par break-even EV / Class B attachment',[f'=SUM(D10:D12)/(1-{inp("cost")})'])
    put(s,19,'EV at historical c69% A reference',[f'=(D10+D11+D12*{inp("a_quote")})/(1-{inp("cost")})'])
    for r in range(10,14):s[f'F{r}'].number_format=PCT
    note(s,22,'Illustrative waterfall, not a legal recovery forecast. Claims, value, costs and ranking are unverified. No cash recovery is added. EV is not RCV. A maturity cash requirement does not remove a claim without funded redemption; the Class A proxy is unchanged.')
    note(s,24,'The c69% Class A reference is reported for March 2026. Applying it to scenario claims is a reverse sensitivity, not an executable price or a March/current market enterprise valuation. No return or buy/sell trade is inferred.')
    s=w['Class B'];head(s,5,['Conditional Class B treatment','Value'])
    rows={6:('Approximate debt written off',f'={inp("b_claim")}'),7:('Implied legacy equity recovery value',f'=D6*{inp("b_recovery")}'),8:('Legacy fraction of assumed post-money equity',f'=D7/{inp("equity")}'),9:('New equity investment paid',f'={inp("b_rights")}'),10:('New investment ownership fraction',f'=D9/{inp("equity")}'),11:('Realised post-money equity value',f'={inp("equity")}*{inp("equity_stress")}'),12:('Marked value of legacy recovery','=D8*D11'),13:('Marked value of new investment','=D10*D11'),14:('Net package value after NEW cash paid','=SUM(D12:D13)-D9'),15:('Legacy recovery rate','=D12/D6'),17:('Equity value for legacy holding to equal old par','=D6/D8')}
    for r,(label,f) in rows.items():put(s,r,label,[f],PCT if r in (8,10,15) else NUM)
    note(s,20,'5.9% is a recovery rate on approximately GBP1bn debt, not a stated ownership percentage. Legacy ownership is derived from an assumed GBP6bn post-money value. Actual equity allocation has not been disclosed here.')
    note(s,22,'The GBP150m right requires fresh investment. It is assumed exercised at fair value. Net package value subtracts that cash; it can be negative when new money loses value. The right itself is not valued. Other 5.9% participation rights and fees are excluded, not treated as free assets.')
    head(s,25,['Equity value multiplier','Legacy value','New value','Net after cash','Legacy recovery'])
    for r,m in enumerate(EQUITY_GRID,26):
        s.cell(r,3,m);s.cell(r,4,f'=$D$7*C{r}');s.cell(r,5,f'=$D$9*C{r}');s.cell(r,6,f'=SUM(D{r}:E{r})-$D$9');s.cell(r,7,f'=D{r}/$D$6')
        for c in range(4,7):s.cell(r,c).number_format=NUM
        s.cell(r,7).number_format=PCT
    note(s,31,'Negotiated Class B equity consideration and a hypothetical enforcement waterfall are separate analyses. They are not cash distributions from the same scenario. No agreed RP2 total equity value, new-money size or complete allocation model is claimed.')
    s=w['Sensitivity'];head(s,5,['Class A recovery: EV / costs',0.,.05,.10])
    for col in 'DEF':s[f'{col}5'].number_format=PCT
    for r,ev in enumerate(EV_GRID,6):
        s.cell(r,3,ev);s.cell(r,3).number_format=NUM
        for col in 'DEF':s[f'{col}{r}']=f'=MAX(0,MIN(1,($C{r}*(1-{col}$5)-Recovery!$D$10-Recovery!$D$11)/Recovery!$D$12))';s[f'{col}{r}'].number_format=PCT
    note(s,14,'Simple live recovery formula with scenario super senior and Class A claims. This tests value and costs, not changes in legal ranking or regulatory outcomes.')
    s.conditional_formatting.add('D6:F11',ColorScaleRule(start_type='num',start_value=0,start_color='F3D5D3',end_type='num',end_value=1,end_color='D5E6D8'))
    s=w['Cover'];head(s,5,['Restructuring & Liability Management','Result'])
    put(s,6,'Call',['REJECT UNCONDITIONAL NEW MONEY']);s.merge_cells('D6:H6')
    for r,label,f,fmt in [(8,'Minimum cash, selected approvals (GBPm)','=Liquidity!D22',NUM),(9,'Peak buffer funding gap (GBPm)','=Liquidity!D23',NUM),(10,'Illustrative Class A waterfall recovery','=Recovery!F12',PCT),(11,'Class A par break-even EV (GBPm)','=Recovery!D18',NUM),(12,'Class B proposed legacy recovery value (GBPm)',"='Class B'!D7",NUM),(13,'Class B legacy ownership at equity proxy',"='Class B'!D8",PCT),(14,'New equity investment is a cash outflow (GBPm)',"='Class B'!D9",NUM)]:put(s,r,label,[f],fmt)
    note(s,17,'Liquidity extension is not a completed recapitalisation. Verify accessible cash, consent conditions, tranche commitments, March maturities, creditor ranking and Ofwat approval before an unconditional commitment. Base funding case assumes approvals; it is not current liquidity.')
    note(s,20,'Do not quote the illustrative waterfall as an actual creditor recovery, or 5.9% as ownership. The model separates proposed consideration for old debt from the cost and risk of participating in new money.')
    s=w['Sources'];head(s,5,['ID','Document / locator','URL']);s.column_dimensions['D'].width=90;s.column_dimensions['E'].width=100
    for r,(key,(desc,url)) in enumerate(SOURCES.items(),6):
        s.cell(r,3,key);s.cell(r,4,desc);s.cell(r,5,url);s.cell(r,5).hyperlink=url;s.cell(r,4).alignment=Alignment(wrap_text=True);s.cell(r,5).alignment=Alignment(wrap_text=True);s.row_dimensions[r].height=75
    for s in w:
        s.print_area=f'C2:{"O" if s.title=="Liquidity" else "K"}{s.max_row}'
        s.sheet_properties.tabColor='1F3864' if s.title=='Cover' else '6E8BB5' if s.title=='Inputs' else None
        for row in s:
            for c in row:
                if c.value is None:continue
                if c.font.name in (None,'Calibri'):
                    colour='000000' if s.title=='Cover' else '008000' if c.data_type=='f' and '!' in c.value else '000000' if c.data_type=='f' else '0000FF' if isinstance(c.value,(int,float)) else '222222'
                    c.font=Font(name='Arial',size=10,color=colour)
                if not c.alignment.wrap_text:c.alignment=Alignment(vertical='center',horizontal='right' if c.data_type=='f' or isinstance(c.value,(int,float)) else 'left')
        s.conditional_formatting.add(f'D6:O{s.max_row}',CellIsRule(operator='lessThan',formula=['0'],font=Font(color='9C0006')))
    with (ROOT/'data/inputs.csv').open('w',newline='') as f:
        writer=csv.writer(f,lineterminator='\n');writer.writerow(['key','label','value','tag','unit','source_or_reason'])
        for k,v in INPUTS.items():writer.writerow([k,*v])
    w.save(RAW);print(RAW)

if __name__=='__main__':main()
