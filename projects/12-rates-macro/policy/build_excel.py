"""Audit workbook: live policy, cash-flow valuation and risk scenarios."""
from pathlib import Path
import json
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter
import assumptions as A
import engine as E
ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'models'/'Gilt_GBP_Policy.xlsx';RAW=OUT.with_name('_raw_'+OUT.name)

def build():
 r=E.results();w=Workbook();w.remove(w.active)
 for n in ['Cover','Inputs','Policy path','Cash flows','Trade','Sensitivity','GBP research','Sources']:w.create_sheet(n)
 s=w['Inputs'];s.append(['Driver','Value','Tag / reason'])
 for row in [('Bank Rate %',A.BANK_RATE,'S: BoE'),('Bank Rate less SONIA bp',A.BASIS_BP,'A: constant basis'),('Gross DV01 GBP/bp',A.BUDGET,'A: sum of absolute legs'),('Horizon years',A.HORIZON,'A: 3 months'),('Scenario terminal Bank Rate %',A.SCENARIO_END,'A: hold, inflation above target'),('Stop adverse bp',A.STOP_BP,'A: risk limit'),('Target favourable bp',A.TARGET_BP,'A: risk/reward'),('OIS 1y forward %',float(E.curves().loc[A.ASOF].ois_fwd_1),'S: BoE continuous forward'),('Gilt 2y spot %',float(E.curves().loc[A.ASOF].gilt_spot_2),'S: BoE fitted zero'),('Gilt 10y spot %',float(E.curves().loc[A.ASOF].gilt_spot_10),'S: BoE fitted zero')]:s.append(row)
 s.append(['Snapshot date',A.ASOF,'S: BoE; FX latest 23 Oct 2026'])
 s.append(['Curve-priced proxies, not quoted bond yields','ISIN','Coupon %','Maturity'])
 for b in A.BONDS:s.append(['S: DMO Annex A p7',*b])
 s=w['Policy path'];s.append(['MPC start','End','t1 years','t2 years','Zero at t1 % D','Zero at t2 % D','Continuous interval %','Implied Bank Rate %','Scenario %'])
 for i,m in enumerate(r['meetings'],2):
  s.append([m['date'],m['end'],m['t1'],m['t2'],E.zero('ois',m['t1'])*100,E.zero('ois',m['t2'])*100,f'=(D{i}*F{i}-C{i}*E{i})/(D{i}-C{i})',f'=365*(EXP(G{i}/100/365)-1)*100+Inputs!B3/100','=Inputs!B6'])
 s.append(['D: linear zero interpolation. Piecewise meeting approximation; risk-neutral plus premia, not a forecast.'])
 s=w['Cash flows'];s.append(['Leg','Time years D','Cash per100 S/D','Zero % D','Aged zero % D','PV','PV -1bp','PV +1bp','Carry PV','Roll PV','Received by horizon'])
 ranges=[]
 for j,b in enumerate(r['bonds']):
  first=s.max_row+1
  for t,c in b['cf']:
   i=s.max_row+1;s.append([j+1,t,c,E.zero('gilt',t)*100,E.zero('gilt',max(t-A.HORIZON,0))*100,f'=C{i}*EXP(-D{i}/100*B{i})',f'=C{i}*EXP(-(D{i}/100-0.0001)*B{i})',f'=C{i}*EXP(-(D{i}/100+0.0001)*B{i})',f'=IF(B{i}>Inputs!B5,C{i}*EXP(-D{i}/100*(B{i}-Inputs!B5)),0)',f'=IF(B{i}>Inputs!B5,C{i}*EXP(-E{i}/100*(B{i}-Inputs!B5)),0)',f'=IF(B{i}<=Inputs!B5,C{i},0)'])
  ranges.append((first,s.max_row))
 s=w['Trade'];s.append(['Metric','Value','Unit / convention'])
 rows=[('12m implied Bank Rate','=365*(EXP(Inputs!B9/100/365)-1)*100+Inputs!B3/100','%'),('Implied cuts','=(Inputs!B2-B2)*100','bp; negative = hikes'),('Scenario cuts','=(Inputs!B2-Inputs!B6)*100','bp'),('2s10s entry','=(Inputs!B11-Inputs!B10)*100','bp fitted zero spread'),('Direction','=IF(Inputs!B6<B2,1,-1)','1 steepener; -1 flattener'),('Stop','=B5-B6*Inputs!B7','bp'),('Target','=B5+B6*Inputs!B8','bp')]
 for row in rows:s.append(row)
 s.append(['Leg','Price /100','DV01 /100','Signed nominal GBP','Carry /100 net funding','Roll /100'])
 for j,(first,last) in enumerate(ranges):
  i=10+j;g=lambda col:f"SUM('Cash flows'!{col}{first}:{col}{last})"
  s.append([A.BONDS[j][0],f'={g("F")}',f'=({g("G")}-{g("H")})/2',f'=Inputs!B4/2/C{i}*100*$B$6*{1 if j==0 else -1}',f'={g("I")}+{g("K")}-B{i}-B{i}*Inputs!B2/100*Inputs!B5',f'={g("J")}-{g("I")}'])
 for row in [('Carry GBP','=D10/100*E10+D11/100*E11','net financing at Bank Rate both legs'),('Roll GBP','=D10/100*F10+D11/100*F11','unchanged spot curve'),('Carry + roll bp','=(B12+B13)/(Inputs!B4/2)','spread-equivalent bp, one leg DV01'),('Net parallel DV01','=D10/100*C10+D11/100*C11','GBP/bp'),('Net cash invested','=D10/100*B10+D11/100*B11','GBP, financed; short proceeds usable assumed')]:s.append(row)
 s=w['Sensitivity'];s.append(['Spread move bp','Linear P&L GBP','Carry + roll GBP','Total GBP'])
 for i,shock in enumerate([-50,-25,0,25,50],2):s.append([shock,f'=A{i}*Trade!B6*Inputs!B4/2','=Trade!B12+Trade!B13',f'=B{i}+C{i}'])
 s.append(['LDI replay','2y shock bp','10y shock bp','Full reprice GBP; no carry'])
 for name,v in r['replay'].items():s.append([name,*v['shocks_bp'],v['pnl_gbp']])
 s.append(['Replay prices current bonds under historical tenor shocks; not a 2022 investable return.'])
 s=w['GBP research'];s.append(['Pair / sample','Slope % per10bp D','HAC t D','R-squared D','N'])
 for pair,v in r['fx'].items():
  for sample,vv in v['regressions'].items():s.append([pair+' '+sample,vv['slope'],vv['t'],vv['r2'],vv['n']])
 s.append(['Carry snapshot','Foreign overnight % S','Annual vol % D','GBP carry %','Carry / vol'])
 row=s.max_row
 panel=E.fx_panel()
 for pair,v in r['fx'].items():
  i=s.max_row+1;s.append([pair,float(panel.iloc[-1].sofr if pair=='GBP_USD' else panel.iloc[-1].estr),v['vol_pct'],f'=Inputs!B2-B{i}',f'=D{i}/C{i}'])
 s.append(['Historical regressions fixed Python research outputs; rerun engine after data refresh.'])
 s.append(['MPC event calendar: 2025-26; CPI: Feb-Oct 2026, n=9; descriptive only.'])
 s.append(['GBP_EUR is GBP appreciation, inverse EUR/GBP. No causal or executable forward-return claim.'])
 s=w['Cover'];s.append(['UK gilts / GBP','27 October 2026'])
 for row in [('Trade',r['trade']),('12m priced cuts bp','=Trade!B3'),('3m carry + roll bp','=Trade!B14'),('Stop fitted spread bp','=Trade!B7'),('Target fitted spread bp','=Trade!B8'),('Gross DV01 GBP/bp','=Inputs!B4'),('Call','Conditional steepener; reject passive GBP carry'),('Limit','Funding, term premia and global FX shocks can overwhelm the view')]:s.append(row)
 s=w['Sources'];s.append(['Source','URL / method'])
 for key,url in A.SOURCES.items():s.append([key,url])
 s.append(['Data audit','policy/data/manifest.json: archive hashes, extraction and retrieval date'])
 s.append(['Curve interpolation D','Linear zero rates; flat below shortest maturity; ACT/365 proxy for gilts'])
 s.append(['Carry A','Symmetric funding at Bank Rate; no bid/ask, repo specialness, margin or haircuts'])
 s.append(['Scenario A','2y -50bp / 10y -10bp gives +40bp spread; inflation / fiscal shock breaks this'])
 for s in w:
  s.freeze_panes='B2';s.sheet_view.showGridLines=False;s.auto_filter.ref=s.dimensions
  for cell in s[1]:cell.fill=PatternFill('solid',fgColor='17365D');cell.font=Font(name='Arial',size=11,bold=True,color='FFFFFF')
  for row in s.iter_rows(min_row=2):
   for c in row:
    c.font=Font(name='Arial',size=10,color='008000' if c.data_type=='f' else '0000FF' if isinstance(c.value,(int,float)) else '222222');c.alignment=Alignment(vertical='top',wrap_text=True)
    if isinstance(c.value,(int,float)) or c.data_type=='f':c.number_format='#,##0.000;[Red](#,##0.000);"—"'
    if c.row%2==0:c.fill=PatternFill('solid',fgColor='EFF3F8')
   s.row_dimensions[row[0].row].height=31
  for j in range(1,s.max_column+1):s.column_dimensions[get_column_letter(j)].width=34 if j==1 else 24
  if s.title in ['Sources','Cover']:s.column_dimensions['B'].width=94
 w.save(RAW);print(RAW)
if __name__=='__main__':build()
